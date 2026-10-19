import pandas as pd
import pytest
from openpyxl import load_workbook

from r2.verify import evaluable as ev
from r2.verify import q3_adjudication as qa


def test_evaluable_pairs_and_discordant_q3():
    p = ev.load_pairs()
    assert len(p) == 333
    assert int(ev.evaluable_mask(p).sum()) == 25
    d = ev.discordant_q3(p)
    assert len(d) == 19
    assert list(d.index) == sorted(d.index)
    assert (d[f"AG_{ev.Q3}"] != d[f"MO_{ev.Q3}"]).all()


def _filled(tmp_path, answers=None, edit=None, drop=None, reverse=False):
    x = tmp_path / "q3.xlsx"
    n = qa.write_workbook(x)
    wb = load_workbook(x)
    ws = wb[qa.SHEET]
    col = qa.COLUMNS.index(f"consensus_{ev.Q3}") + 1
    for r in range(2, n + 2):
        ws.cell(row=r, column=col).value = (answers or {}).get(r, "y ")
    if edit:
        r, name, value = edit
        ws.cell(row=r, column=qa.COLUMNS.index(name) + 1).value = value
    if drop:
        ws.delete_rows(drop)
    if reverse:
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2, max_row=ws.max_row)]
        for i, vals in enumerate(reversed(rows), start=2):
            for j, v in enumerate(vals, start=1):
                ws.cell(row=i, column=j).value = v
    wb.save(x)
    return x, n


def test_workbook_rows_and_round_trip(tmp_path):
    x, n = _filled(tmp_path, reverse=True)
    assert n == 19
    out = qa.import_workbook(x, "2026-10-20", csv_path=tmp_path / "c.csv")
    assert len(out) == 19 and set(out[f"consensus_{ev.Q3}"]) == {"Y"}      # "y " normalized
    assert list(out.case_id) == sorted(out.case_id)                         # matched by case_id, not row order
    assert (pd.read_csv(tmp_path / "c.csv").adjudicated_on == "2026-10-20").all()


@pytest.mark.parametrize("kind", ["blank", "unknown", "changed_original", "missing_row"])
def test_import_refuses(tmp_path, kind):
    kw = {"blank": {"answers": {2: None}}, "unknown": {"answers": {3: "maybe"}},
          "changed_original": {"edit": (4, f"AG_{ev.Q3}", "Y")}, "missing_row": {"drop": 5}}[kind]
    x, _ = _filled(tmp_path, **kw)
    with pytest.raises(ValueError):
        qa.import_workbook(x, "2026-10-20", csv_path=tmp_path / "c.csv")
    assert not (tmp_path / "c.csv").exists()


def test_import_refuses_bad_date(tmp_path):
    x, _ = _filled(tmp_path)
    with pytest.raises(ValueError):
        qa.import_workbook(x, "20 Oct", csv_path=tmp_path / "c.csv")
