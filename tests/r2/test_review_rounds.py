"""Round-based review packets, Excel workbooks, author-edit import and round merging (tmp_path only)."""
import csv
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from r2.verify import build_packets as bp
from r2.verify import score_review as sr

QKEYS = [q for q, _ in sr.QUESTIONS]


@pytest.fixture
def env(tmp_path, monkeypatch, make_pairs):
    data, verify = tmp_path / "data", tmp_path / "verification"
    data.mkdir()
    monkeypatch.setattr(bp, "DATA_V2", data)
    monkeypatch.setattr(bp, "VERIFY_DIR", verify)
    monkeypatch.setattr(sr, "VERIFY_DIR", verify)
    (data / "pairs_v2.json").write_text(json.dumps(make_pairs(["C1", "C2", "C3"], unresolved=["U1", "U2"])))
    return data, verify


def _queue(path, ids=("U1", "U2")):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["case_id", "sham_type", "authentic_organization", "clinical_question", "authentic_excerpt_text"])
        for i, cid in enumerate(ids):
            w.writerow([cid, "subtle_inversion", "Org", f"Q{i}?", "Use drug A. Avoid drug B in renal failure."])
    return path


def test_round1_forms_resolved_only(env):
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    d = verify / "round1"
    for name in ("review_AG.csv", "review_MO.csv"):
        assert set(pd.read_csv(d / name)["case_id"]) == {"C1", "C2", "C3"}
    for name in ("packets.html", "README_REVIEWERS.md", "packets_manifest.json"):
        assert (d / name).exists()
    html = (d / "packets.html").read_text()
    assert "U1" not in html and "U2" not in html
    m = json.loads((d / "packets_manifest.json").read_text())
    assert m["unresolved"] == ["U1", "U2"] and set(m["packets"]) == {"C1", "C2", "C3"}
    assert not (verify / "review_AG.csv").exists()


def test_round1_fixed_shuffles_differ_and_are_stable(env):
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    first = (verify / "round1" / "review_AG.csv").read_bytes()
    bp.write_forms("round1", allow_unresolved=True)
    assert (verify / "round1" / "review_AG.csv").read_bytes() == first


def test_default_write_forms_still_refuses_unresolved(env):
    data, verify = env
    with pytest.raises(RuntimeError):
        bp.write_forms()
    assert not verify.exists()


def test_round_forms_refuse_overwrite_of_answers(env):
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    p = verify / "round1" / "review_MO.csv"
    df = pd.read_csv(p, dtype=str).fillna("")
    df.loc[0, "q1_vignette_consistent"] = "Y"
    df.to_csv(p, index=False)
    with pytest.raises(RuntimeError, match="review_MO.csv"):
        bp.write_forms("round1", allow_unresolved=True)


def test_write_round_forms_subset(env):
    data, verify = env
    bp.write_round_forms("sub", ["C2"])
    d = verify / "sub"
    assert list(pd.read_csv(d / "review_AG.csv")["case_id"]) == ["C2"]
    m = json.loads((d / "packets_manifest.json").read_text())
    assert set(m["packets"]) == {"C2"} and m["case_ids"] == ["C2"]
    with pytest.raises(ValueError):
        bp.write_round_forms("sub2", ["NOPE"])


def test_xlsx_roundtrip_exact(env, tmp_path):
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    src = verify / "round1" / "review_AG.csv"
    x = tmp_path / "ag.xlsx"
    bp.export_xlsx(src, x, "A.G.", "round1")
    out = tmp_path / "back.csv"
    bp.import_review_xlsx(x, out)
    assert out.read_bytes() == src.read_bytes()


def test_xlsx_sheets_and_dropdowns(env, tmp_path):
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    x = tmp_path / "mo.xlsx"
    bp.export_xlsx(verify / "round1" / "review_MO.csv", x, "M.O.", "round1")
    wb = load_workbook(x)
    assert wb.sheetnames == ["Instructions", "Review"]
    text = "\n".join(str(c.value) for row in wb["Instructions"].iter_rows() for c in row if c.value)
    assert "M.O." in text and "do not discuss until both are returned" in text
    ws = wb["Review"]
    header = [c.value for c in ws[1]]
    assert header == list(pd.read_csv(verify / "round1" / "review_MO.csv", nrows=0).columns)
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    got = {}
    for dv in ws.data_validations.dataValidation:
        col = str(dv.sqref)[0]
        assert dv.type == "list"
        got[header[ord(col) - 65]] = dv.formula1
    assert set(got) == set(QKEYS)
    for q in QKEYS:
        assert got[q] == ('"Y,N,P"' if q == "q3_authentic_concordant" else '"Y,N"')


def test_ag_workbook_has_author_edits_mo_does_not(env, tmp_path):
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    q = _queue(tmp_path / "q.csv", [f"U{i}" for i in range(23)])
    ag, mo = tmp_path / "ag.xlsx", tmp_path / "mo.xlsx"
    bp.export_xlsx(verify / "round1" / "review_AG.csv", ag, "A.G.", "round1", queue_csv=q)
    bp.export_xlsx(verify / "round1" / "review_MO.csv", mo, "M.O.", "round1")
    assert "Author_edits" not in load_workbook(mo).sheetnames
    ws = load_workbook(ag)["Author_edits"]
    assert "ONE minimal edit" in str(ws["A1"].value)
    hdr = [c.value for c in ws[2]]
    assert hdr == ["case_id", "sham_type", "clinical_question", "authentic_organization",
                   "authentic_excerpt_text", "field", "find", "replace", "rationale"]
    assert [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)] == [f"U{i}" for i in range(23)]
    assert any(dv.formula1 == '"excerpt_text,metadata.organization"' for dv in ws.data_validations.dataValidation)


def _author_sheet(tmp_path, monkeypatch, rows):
    from openpyxl import Workbook
    from r2.stimuli import transforms  # noqa: F401
    orig = {"U1": {"eval": {"sham_trap_type": "subtle_inversion"},
                   "tools": {"S": {"metadata": {"organization": "Org"},
                                   "excerpt_text": "Use drug A. Avoid drug B. Avoid drug B2."}}}}
    monkeypatch.setattr(bp, "load_original_cases", lambda: (orig, {}))
    wb = Workbook()
    ws = wb.active
    ws.title = "Author_edits"
    ws.append(["instructions"])
    ws.append(["case_id", "sham_type", "clinical_question", "authentic_organization",
               "authentic_excerpt_text", "field", "find", "replace", "rationale"])
    for r in rows:
        ws.append(["U1", "subtle_inversion", "", "", "", *r])
    p = tmp_path / "ae.xlsx"
    wb.save(p)
    return p


def test_import_author_edits_accepts_valid(env, tmp_path, monkeypatch):
    data, _ = env
    p = _author_sheet(tmp_path, monkeypatch, [("excerpt_text", "Use drug A", "Do not use drug A", "flip")])
    out = bp.import_author_edits(p, "A.G.", "2026-10-06")
    assert out == {"U1": {"field": "excerpt_text", "find": "Use drug A", "replace": "Do not use drug A",
                          "author": "A.G.", "date": "2026-10-06", "rationale": "flip"}}
    assert json.loads((data / "manual_edits.json").read_text()) == out


@pytest.mark.parametrize("row", [
    ("excerpt_text", "Avoid drug B", "Use drug B", ""),   # occurs twice
    ("excerpt_text", "Use drug A", "", ""),               # blank replace
    ("excerpt_text", "", "x", ""),                         # blank find
    ("bogus", "Use drug A", "x", ""),                     # unknown field
])
def test_import_author_edits_rejects(env, tmp_path, monkeypatch, row):
    data, _ = env
    p = _author_sheet(tmp_path, monkeypatch, [row])
    with pytest.raises(ValueError, match="U1"):
        bp.import_author_edits(p, "A.G.", "2026-10-06")
    assert not (data / "manual_edits.json").exists()


def _answer(path):
    df = pd.read_csv(path, dtype=str).fillna("")
    for q in QKEYS:
        df[q] = "N" if q == "q8_leakage" else "Y"
    df.to_csv(path, index=False)


def _two_rounds(env, make_pairs):
    data, verify = env
    (data / "pairs_v2.json").write_text(json.dumps(make_pairs(["C1", "C2"], unresolved=["C3"])))
    bp.write_forms("round1", allow_unresolved=True)
    (data / "pairs_v2.json").write_text(json.dumps(make_pairs(["C1", "C2", "C3"])))
    bp.write_round_forms("round2", ["C3"])
    for rd in ("round1", "round2"):
        for r in ("AG", "MO"):
            _answer(verify / rd / f"review_{r}.csv")
    return data, verify


def test_merge_feeds_apply_consensus(env, make_pairs):
    data, verify = _two_rounds(env, make_pairs)
    # round1 manifest still names C1,C2 only (C3 unresolved then)
    bp.merge_rounds(["round1", "round2"])
    a = pd.read_csv(verify / "review_AG.csv", dtype=str)
    assert sorted(a.case_id) == ["C1", "C2", "C3"]
    m = json.loads((verify / "packets_manifest.json").read_text())
    assert set(m["packets"]) == {"C1", "C2", "C3"}
    cons = verify / "c.csv"
    pd.DataFrame([], columns=["case_id", "item", "consensus"]).to_csv(cons, index=False)
    excluded, _ = sr.apply_consensus(verify / "review_AG.csv", verify / "review_MO.csv", cons)
    assert excluded == {}


def test_merge_refuses_duplicate_case(env, make_pairs):
    data, verify = _two_rounds(env, make_pairs)
    bp.write_round_forms("round3", ["C1"])
    _answer(verify / "round3" / "review_AG.csv")
    _answer(verify / "round3" / "review_MO.csv")
    with pytest.raises(ValueError, match="C1"):
        bp.merge_rounds(["round1", "round3"])
    assert not (verify / "review_AG.csv").exists()


def test_merge_refuses_stale_packet(env, make_pairs):
    data, verify = _two_rounds(env, make_pairs)
    pairs = json.loads((data / "pairs_v2.json").read_text())
    pairs["cases"][0]["tools"]["S"]["excerpt_text"] = "Give drug 0 at 99 mg daily."
    (data / "pairs_v2.json").write_text(json.dumps(pairs))
    with pytest.raises(ValueError, match="packet_sha256"):
        bp.merge_rounds(["round1", "round2"])
    assert not (verify / "review_AG.csv").exists()


# ---- Task DM 7b minors ----
def test_import_author_edits_accepts_only_queued_case_ids(env, tmp_path, monkeypatch):
    from openpyxl import Workbook
    data, _ = env  # pairs_v2.json queues U1 and U2
    orig = {cid: {"eval": {"sham_trap_type": "subtle_inversion"},
                  "tools": {"S": {"metadata": {"organization": "Org"}, "excerpt_text": "Use drug A."}}}
            for cid in ("U1", "C1")}
    monkeypatch.setattr(bp, "load_original_cases", lambda: (orig, {}))
    wb = Workbook()
    ws = wb.active
    ws.title = "Author_edits"
    ws.append(["case_id", "sham_type", "clinical_question", "authentic_organization",
               "authentic_excerpt_text", "field", "find", "replace", "rationale"])
    for cid in ("U1", "C1"):
        ws.append([cid, "", "", "", "", "excerpt_text", "Use drug A", "Do not use drug A", ""])
    p = tmp_path / "ae.xlsx"
    wb.save(p)
    with pytest.raises(ValueError, match="C1.*not in the manual-edit queue"):
        bp.import_author_edits(p, "A.G.", "2026-10-06")
    assert not (data / "manual_edits.json").exists()


def test_merge_requires_full_coverage_early(env, make_pairs):
    data, verify = env
    (data / "pairs_v2.json").write_text(json.dumps(make_pairs(["C1", "C2"], unresolved=["C3"])))
    bp.write_forms("round1", allow_unresolved=True)
    (data / "pairs_v2.json").write_text(json.dumps(make_pairs(["C1", "C2", "C3"])))
    for r in ("AG", "MO"):
        _answer(verify / "round1" / f"review_{r}.csv")
    with pytest.raises(ValueError, match="do not cover.*C3"):
        bp.merge_rounds(["round1"])
    assert not (verify / "review_AG.csv").exists() and not (verify / "packets_manifest.json").exists()


def _round1_xlsx(env, tmp_path):
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    x = tmp_path / "ag.xlsx"
    bp.export_xlsx(verify / "round1" / "review_AG.csv", x, "A.G.", "round1")
    return x


def _edit_cell(x, column, row_idx, value):
    wb = load_workbook(x)
    ws = wb["Review"]
    col = [c.value for c in ws[1]].index(column) + 1
    ws.cell(row_idx, col).value = value
    wb.save(x)


@pytest.mark.parametrize("column", ["vignette_as_shown_to_model", "modified_as_shown_to_model",
                                    "packet_sha256", "case_id", "clinical_question"])
def test_import_review_refuses_edited_content(env, tmp_path, column):
    x = _round1_xlsx(env, tmp_path)
    _edit_cell(x, column, 2, "tampered")
    out = tmp_path / "back.csv"
    with pytest.raises(ValueError, match="content|case_id"):
        bp.import_review_xlsx(x, out)
    assert not out.exists()


def test_import_review_accepts_answers_and_notes(env, tmp_path):
    x = _round1_xlsx(env, tmp_path)
    _edit_cell(x, "q1_vignette_consistent", 2, "Y")
    _edit_cell(x, "notes", 2, "looks fine")
    out = tmp_path / "back.csv"
    bp.import_review_xlsx(x, out)
    assert "looks fine" in out.read_text(encoding="utf-8")
