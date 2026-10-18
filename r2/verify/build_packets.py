"""Physician review packets showing exactly what models receive (rendered by r2.stimuli.render)."""
import csv
import difflib
import hashlib
import html
import json
import random
import re
from pathlib import Path

from r2.io import load_original_cases
from r2.paths import DATA_V2, VERIFY_DIR
from r2.stimuli.render import render_tool, render_vignette
from r2.stimuli.transforms import TransformError, apply_manual_edit
from r2.verify.score_review import QUESTIONS

# The packet hash binds each physician-reviewed row to the content the models receive.
PACKET_FIELDS = ("vignette_as_shown_to_model", "clinical_question", "authentic_as_shown_to_model",
                 "modified_as_shown_to_model")
PACKET_SEP = "\n\x1e\n"
REVIEWERS = (("AG", 1), ("MO", 2))
REVIEWER_NAMES = {"AG": "A.G.", "MO": "M.O."}  # as export_xlsx writes them into Instructions!A1
MANIFEST = "packets_manifest.json"


def pairs_path():
    return DATA_V2 / "pairs_v2.json"


def _load():
    return json.loads(pairs_path().read_text(encoding="utf-8"))


def _cases():
    return _load()["cases"]


def packet_sha256(row):
    parts = ["" if row[k] is None else str(row[k]) for k in PACKET_FIELDS]
    return hashlib.sha256(PACKET_SEP.join(parts).encode("utf-8")).hexdigest()


def packet_rows(cases=None):
    rows = []
    for c in (_cases() if cases is None else cases):
        t, s = c["tools"]["T"], c["tools"]["S"]
        rec = c["edit_record"]
        row = {
            "case_id": c["case_id"], "sham_type": c["eval"]["sham_trap_type"],
            "source_guideline": f'{t["metadata"].get("organization")} - {t["metadata"].get("guideline_title")} '
                                f'({t["metadata"].get("publication_year")}); file '
                                f'{Path(t["metadata"].get("source_pdf_path") or "").name}',
            "vignette_as_shown_to_model": render_vignette(c),
            "clinical_question": c.get("clinical_question"),
            "authentic_as_shown_to_model": json.dumps(render_tool(t, "authentic"), indent=1, ensure_ascii=False),
            "modified_as_shown_to_model": json.dumps(render_tool(s, "modified"), indent=1, ensure_ascii=False),
            "what_changed": json.dumps(rec["changes"], ensure_ascii=False),
            "edit_source": rec["edit_source"],
        }
        row["packet_sha256"] = packet_sha256(row)
        rows.append(row)
    return rows


def _answer_cols():
    return [q for q, _ in QUESTIONS] + ["notes"]


def _answered_forms(answer_cols, directory=None):
    directory = VERIFY_DIR if directory is None else directory
    answered = []
    for reviewer, _ in REVIEWERS:
        path = directory / f"review_{reviewer}.csv"
        if not path.exists():
            continue
        with open(path, newline="", encoding="utf-8") as fh:
            if any((row.get(c) or "").strip() for row in csv.DictReader(fh) for c in answer_cols):
                answered.append(path.name)
    return answered


def readme_text():
    guide = ["# Reviewer instructions", "",
             "Review every row independently; do not discuss until both forms are complete.",
             "The *_as_shown_to_model columns are exactly what the evaluated models receive, except for the first line: here it reads \"label\": \"authentic\" or \"modified\" so that you know which is which; the models see \"label\": \"A\" or \"B\", and every pair is shown in both orders.",
             "Answer each item with Y or N (Q3 also allows P = partially). Use `notes` for any N.", ""]
    guide += [f"- **{q}**: {text}" for q, text in QUESTIONS]
    guide += ["",
              "Edits produced by rule (edit_source = rule_typosquat or rule_inversion_extended) were generated automatically and may read awkwardly (e.g. an ungrammatical or clinically odd insertion). If an edit is clinically meaningful but contains a grammar or capitalization artifact that a model could use as a surface cue, answer Q6 = N and write a corrected edit in notes.",
              "",
              "Some authentic excerpts were selected by clinical domain, not by question, and may not address the clinical question (e.g. CASE_133). Answer Q4 accordingly.",
              "",
              "CASE_335 has no clinical question (rendered 'Not specified'); judge Q2 and Q4 accordingly.",
              "",
              "18 of the 50 allergy_ignorance cases have no recorded allergy (rendered 'No known drug allergies'); check Q7 (does the modification reduce trust for this patient?) explicitly for these.",
              "",
              "Cases listed in verification/leakage_flags.json contain allow-listed substrings (e.g. 'gi_bleed' in applicability text); check them under Q8."]
    return "\n".join(guide) + "\n"


def _write_forms_to(directory, raw, cases, extra_manifest=None):
    answer_cols = _answer_cols()
    answered = _answered_forms(answer_cols, directory)
    if answered:
        raise RuntimeError(f"refusing to overwrite review forms that already contain answers: {answered}")
    directory.mkdir(parents=True, exist_ok=True)
    rows = packet_rows(cases)
    manifest = {"pairs_v2_sha256": hashlib.sha256(raw).hexdigest(),
                "packets": {r["case_id"]: r["packet_sha256"] for r in rows}}
    manifest.update(extra_manifest or {})
    (directory / MANIFEST).write_text(json.dumps(manifest, indent=1) + "\n", encoding="utf-8")
    for reviewer, seed in REVIEWERS:
        order = rows[:]
        random.Random(seed).shuffle(order)
        with open(directory / f"review_{reviewer}.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0]) + answer_cols)
            w.writeheader()
            for r in order:
                w.writerow({**r, **{c: "" for c in answer_cols}})
    differ = difflib.HtmlDiff(wrapcolumn=80)
    parts = ["<html><meta charset='utf-8'><body style='font-family:Arial'>"]
    for r in rows:
        parts.append(f"<h2 id='{r['case_id']}'>{r['case_id']} ({r['sham_type']})</h2>")
        parts.append(f"<pre>{html.escape(r['vignette_as_shown_to_model'])}</pre>")
        parts.append(f"<p><b>Question:</b> {html.escape(str(r['clinical_question']))}</p>")
        parts.append(differ.make_table(r["authentic_as_shown_to_model"].splitlines(),
                                       r["modified_as_shown_to_model"].splitlines(),
                                       "Authentic", "Modified"))
    parts.append("</body></html>")
    (directory / "packets.html").write_text("\n".join(parts), encoding="utf-8")
    (directory / "README_REVIEWERS.md").write_text(readme_text(), encoding="utf-8")


def write_forms(round_label="round1", allow_unresolved=False):
    """Default: the full 500-pair forms in verification/ (refuses unless complete). With
    allow_unresolved=True: forms for the resolved cases only, in verification/<round_label>/."""
    raw = pairs_path().read_bytes()
    d = json.loads(raw)
    if not allow_unresolved:
        if len(d["cases"]) != 500 or d["unresolved"]:
            raise RuntimeError(
                f"refusing to write review forms: pairs_v2.json has {len(d['cases'])} cases and "
                f"{len(d['unresolved'])} unresolved pairs; need exactly 500 cases and none unresolved "
                "(physician-authored manual edits must be applied and the build rerun strictly)")
        _write_forms_to(VERIFY_DIR, raw, d["cases"])
        return
    _write_forms_to(VERIFY_DIR / round_label, raw, d["cases"], {"unresolved": list(d["unresolved"])})


def write_round_forms(round_label, case_ids):
    """Forms for an explicit subset of the current pairs (e.g. round 2)."""
    raw = pairs_path().read_bytes()
    by_id = {c["case_id"]: c for c in json.loads(raw)["cases"]}
    missing = [cid for cid in case_ids if cid not in by_id]
    if missing:
        raise ValueError(f"case_ids not resolved in pairs_v2.json: {missing}")
    _write_forms_to(VERIFY_DIR / round_label, raw, [by_id[cid] for cid in case_ids],
                    {"case_ids": list(case_ids)})


# ---- Excel workbooks ----
INSTRUCTIONS = "Instructions"
REVIEW = "Review"
AUTHOR_EDITS = "Author_edits"
AUTHOR_COLS = ["case_id", "sham_type", "clinical_question", "authentic_organization",
               "authentic_excerpt_text", "field", "find", "replace", "rationale"]
EDIT_FIELDS = ("excerpt_text", "metadata.organization")
AUTHOR_INSTRUCTION = (
    "Write ONE minimal edit per case, of the declared type. subtle_inversion: invert the polarity of one "
    "recommendation in excerpt_text (field = excerpt_text). authority_mimicry: a plausible look-alike of the "
    "organization name (field = metadata.organization). `find` must be copied exactly from the authentic text "
    "(authentic_excerpt_text or authentic_organization) and occur exactly once; `replace` is what it becomes. "
    "Avoid grammar or capitalization artifacts a model could use as a surface cue. Optional: rationale.")
LONG_COLS = {"vignette_as_shown_to_model", "authentic_as_shown_to_model", "modified_as_shown_to_model",
             "what_changed", "source_guideline", "clinical_question", "notes", "authentic_excerpt_text",
             "find", "replace", "rationale"}


def _set(ws, row, col, value):
    from openpyxl.cell.cell import Cell  # noqa: F401
    c = ws.cell(row=row, column=col)
    c.value = value
    if isinstance(value, str):
        c.data_type = "s"  # never let a leading "=" turn text into a formula
    return c


def _dropdown(ws, col_letter, first, last, values):
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Invalid", error=f"Choose one of: {', '.join(values)}")
    dv.add(f"{col_letter}{first}:{col_letter}{max(last, first)}")
    ws.add_data_validation(dv)


def _style_table(ws, header_row, columns):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    for i, name in enumerate(columns, start=1):
        h = ws.cell(row=header_row, column=i)
        h.font = Font(bold=True)
        h.fill = PatternFill("solid", fgColor="D9E1F2")
        letter = get_column_letter(i)
        if name in LONG_COLS:
            ws.column_dimensions[letter].width = 60 if name != "notes" else 40
        elif name == "packet_sha256":
            ws.column_dimensions[letter].width = 18
        else:
            ws.column_dimensions[letter].width = 16 if len(name) < 16 else min(len(name) + 2, 28)
        for r in range(header_row + 1, ws.max_row + 1):
            ws.cell(row=r, column=i).alignment = Alignment(wrap_text=name in LONG_COLS, vertical="top")
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{max(ws.max_row, header_row + 1)}"


def export_xlsx(csv_path, xlsx_path, reviewer, round_label, queue_csv=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    with open(csv_path, newline="", encoding="utf-8") as fh:
        table = list(csv.reader(fh))
    header, body = table[0], table[1:]
    wb = Workbook()
    ins = wb.active
    ins.title = INSTRUCTIONS
    ins.column_dimensions["A"].width = 130
    lines = [f"Reviewer: {reviewer}    Round: {round_label}",
             "Answer independently; do not discuss until both are returned.",
             "Fill the dropdown answer columns on the Review sheet (Y/N; Q3 also P). Do not edit or reorder other columns.",
             ""] + readme_text().splitlines()
    if queue_csv is not None:
        lines += ["", f"Sheet {AUTHOR_EDITS}: " + AUTHOR_INSTRUCTION]
    for i, line in enumerate(lines, start=1):
        c = _set(ins, i, 1, line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ins["A1"].font = Font(bold=True, size=13)
    ws = wb.create_sheet(REVIEW)
    for j, name in enumerate(header, start=1):
        _set(ws, 1, j, name)
    for i, row in enumerate(body, start=2):
        for j, val in enumerate(row, start=1):
            if val != "":
                _set(ws, i, j, val)
    _style_table(ws, 1, header)
    for q, _ in QUESTIONS:
        if q in header:
            _dropdown(ws, get_column_letter(header.index(q) + 1), 2, len(body) + 1,
                      ("Y", "N", "P") if q == "q3_authentic_concordant" else ("Y", "N"))
    if queue_csv is not None:
        with open(queue_csv, newline="", encoding="utf-8") as fh:
            queue = list(csv.DictReader(fh))
        ae = wb.create_sheet(AUTHOR_EDITS)
        c = _set(ae, 1, 1, AUTHOR_INSTRUCTION)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ae.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(AUTHOR_COLS))
        ae.row_dimensions[1].height = 75
        for j, name in enumerate(AUTHOR_COLS, start=1):
            _set(ae, 2, j, name)
        for i, q in enumerate(queue, start=3):
            for j, name in enumerate(AUTHOR_COLS, start=1):
                if q.get(name):
                    _set(ae, i, j, q[name])
        _style_table(ae, 2, AUTHOR_COLS)
        _dropdown(ae, get_column_letter(AUTHOR_COLS.index("field") + 1), 3, len(queue) + 2, EDIT_FIELDS)
    wb.save(xlsx_path)


def _cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def _round_dir(round_label):
    return VERIFY_DIR / round_label


def _check_review_content(rows, xlsx_path, round_label, reviewer):
    """Everything except the answer columns must equal the form as sent in `round_label`: the case_id must be in
    that round's manifest, the packet_sha256 cell must equal the manifest value and the hash recomputed from
    the row's PACKET_FIELDS, and the other content columns must equal that round's form CSV. A reviewer edit to
    a case_id, a packet column or its hash is refused. (The current render is not consulted: a pair may have
    changed since the form was sent.)"""
    if not rows:
        raise ValueError(f"{xlsx_path}: the '{REVIEW}' sheet is empty")
    header = rows[0]
    answer = set(_answer_cols())
    d = _round_dir(round_label)
    manifest = json.loads((d / MANIFEST).read_text(encoding="utf-8"))["packets"]
    with open(d / f"review_{reviewer}.csv", newline="", encoding="utf-8") as fh:
        form_rd = csv.DictReader(fh)
        form_cols = list(form_rd.fieldnames or [])
        form = {r["case_id"]: r for r in form_rd}
    content_cols = [c for c in header if c not in answer]
    expected_cols = [c for c in form_cols if c not in answer]
    if content_cols != expected_cols:
        raise ValueError(f"{xlsx_path}: content columns {content_cols} differ from the packet columns {expected_cols}")
    problems = []
    for line, row in enumerate(rows[1:], start=2):
        cells = dict(zip(header, row))
        cid = cells.get("case_id", "")
        if cid not in manifest or cid not in form:
            problems.append(f"row {line}: case_id {cid!r} is not a case_id of {round_label}")
            continue
        if cells.get("packet_sha256", "") != manifest[cid]:
            problems.append(f"row {line} ({cid}): packet_sha256 differs from the {round_label} manifest")
        elif packet_sha256(cells) != manifest[cid]:
            problems.append(f"row {line} ({cid}): packet content does not match its {round_label} packet_sha256")
        changed = [c for c in content_cols if cells.get(c, "") != form[cid][c]]
        if changed:
            problems.append(f"row {line} ({cid}): content columns differ from the {round_label} form: {changed}")
    if problems:
        raise ValueError(f"{xlsx_path}: a reviewer edited packet content, nothing imported:\n  "
                         + "\n  ".join(problems[:20]))


def _check_row_order(rows, xlsx_path, round_label, reviewer):
    """The Review sheet's case_ids in the order of `reviewer`'s form as sent. The reviewers' forms hold the same
    packets in different fixed shuffles, so the order identifies whose form a workbook is."""
    with open(_round_dir(round_label) / f"review_{reviewer}.csv", newline="", encoding="utf-8") as fh:
        sent = [r["case_id"] for r in csv.DictReader(fh)]
    col = rows[0].index("case_id") if rows and "case_id" in rows[0] else 0
    if [r[col] if col < len(r) else "" for r in rows[1:]] != sent:
        raise ValueError(f"{xlsx_path}: the row order of the '{REVIEW}' sheet is not that of "
                         f"{REVIEWER_NAMES[reviewer]}'s {round_label} form, so the workbook is not "
                         f"{REVIEWER_NAMES[reviewer]}'s")


def import_review_xlsx(xlsx_path, csv_path, round_label, reviewer=None):
    """Write the returned workbook's Review sheet to csv_path after checking it against the forms of
    `round_label` as they were sent (that round's manifest and form CSV).

    The reviewer is read from Instructions!A1. `reviewer` (a code of REVIEWERS) is needed only when that header
    was cleared in the returned copy: the Review sheet's row order must then equal that reviewer's form as sent.
    When the header is present, `reviewer` must agree with it."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=False, data_only=True)
    if REVIEW not in wb.sheetnames:
        raise ValueError(f"{xlsx_path}: no '{REVIEW}' sheet")
    a1 = str(wb[INSTRUCTIONS]["A1"].value or "") if INSTRUCTIONS in wb.sheetnames else ""
    m = re.fullmatch(r"Reviewer: (\S+)\s+Round: (\S+)", a1.strip())
    names = {v: k for k, v in REVIEWER_NAMES.items()}
    if reviewer is not None and reviewer not in REVIEWER_NAMES:
        raise ValueError(f"unknown reviewer {reviewer!r}; expected one of {sorted(REVIEWER_NAMES)}")
    headerless = not a1.strip() and reviewer is not None
    if not headerless:
        if not m or m.group(1) not in names or m.group(2) != round_label:
            hint = " (if the header was cleared on return, name the reviewer with reviewer=)" if not a1.strip() else ""
            raise ValueError(f"{xlsx_path}: {INSTRUCTIONS}!A1 reads {a1!r}, expected a reviewer and Round: "
                             f"{round_label}{hint}")
        if reviewer is not None and names[m.group(1)] != reviewer:
            raise ValueError(f"{xlsx_path}: {INSTRUCTIONS}!A1 names {m.group(1)}, not {REVIEWER_NAMES[reviewer]}")
        reviewer = names[m.group(1)]
    rows = [[_cell_text(c.value) for c in row] for row in wb[REVIEW].iter_rows()]
    while rows and not any(rows[-1]):
        rows.pop()
    if Path(csv_path).name != f"review_{reviewer}.csv":
        raise ValueError(f"{xlsx_path}: the workbook is {REVIEWER_NAMES[reviewer]}'s but the target is "
                         f"{Path(csv_path).name}, expected review_{reviewer}.csv")
    _check_review_content(rows, xlsx_path, round_label, reviewer)
    if headerless:
        _check_row_order(rows, xlsx_path, round_label, reviewer)
    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(rows)


def import_author_edits(xlsx_path, author, date):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, data_only=True)
    if AUTHOR_EDITS not in wb.sheetnames:
        raise ValueError(f"{xlsx_path}: no '{AUTHOR_EDITS}' sheet")
    rows = [[_cell_text(c.value) for c in row] for row in wb[AUTHOR_EDITS].iter_rows()]
    start = next((i for i, r in enumerate(rows) if r and r[0] == "case_id"), None)
    if start is None:
        raise ValueError(f"{AUTHOR_EDITS}: header row (starting with case_id) not found")
    col = {name: rows[start].index(name) for name in AUTHOR_COLS if name in rows[start]}
    originals, _ = load_original_cases()
    queued = set(_load()["unresolved"])
    edits, problems = {}, []
    for r in rows[start + 1:]:
        cid = r[col["case_id"]].strip()
        if not cid:
            continue
        get = lambda n: r[col[n]] if n in col and col[n] < len(r) else ""  # noqa: E731
        edit = {"field": get("field").strip(), "find": get("find"), "replace": get("replace"),
                "author": author, "date": date}
        if get("rationale").strip():
            edit["rationale"] = get("rationale").strip()
        if cid in edits:
            problems.append(f"{cid}: appears more than once")
        elif cid not in originals:
            problems.append(f"{cid}: unknown case_id")
        elif cid not in queued:
            problems.append(f"{cid}: not in the manual-edit queue (pairs_v2.json unresolved list)")
        elif not edit["find"].strip() or not edit["replace"].strip():
            problems.append(f"{cid}: blank find or replace")
        elif edit["field"] not in EDIT_FIELDS:
            problems.append(f"{cid}: unknown field {edit['field']!r}")
        else:
            expected = {"subtle_inversion": "excerpt_text", "authority_mimicry": "metadata.organization"}.get(
                originals[cid]["eval"]["sham_trap_type"])
            try:
                if expected and edit["field"] != expected:
                    raise TransformError(f"{originals[cid]['eval']['sham_trap_type']} edits belong in {expected}")
                apply_manual_edit(originals[cid]["tools"]["S"], edit)
            except (TransformError, KeyError) as e:
                problems.append(f"{cid}: {e}")
            else:
                edits[cid] = edit
    if problems:
        raise ValueError("author edits rejected, nothing written:\n  " + "\n  ".join(problems))
    path = DATA_V2 / "manual_edits.json"
    merged = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    merged.update(edits)
    path.write_text(json.dumps(merged, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    return edits


def merge_rounds(round_labels, out_dir=None):
    """Concatenate the rounds' forms per reviewer into out_dir/review_<R>.csv and write the top-level
    manifest for the current pairs_v2.json. A row from an earlier round whose packet_sha256 differs from the
    current render is superseded, and dropped, when a later round holds that case_id with the current hash
    (the manifest lists these under "superseded"); a case_id with the current hash in two rounds, and a stale
    row nothing supersedes, are errors."""
    out_dir = VERIFY_DIR if out_dir is None else Path(out_dir)
    raw = pairs_path().read_bytes()
    current = {r["case_id"]: r["packet_sha256"] for r in packet_rows(json.loads(raw)["cases"])}
    merged, superseded = {}, {}
    for reviewer, _ in REVIEWERS:
        header, entries = None, []  # (label, row) in round order
        for label in round_labels:
            path = VERIFY_DIR / label / f"review_{reviewer}.csv"
            with open(path, newline="", encoding="utf-8") as fh:
                rd = csv.DictReader(fh)
                if header is None:
                    header = rd.fieldnames
                elif rd.fieldnames != header:
                    raise ValueError(f"{path}: columns differ from the first round")
                entries += [(label, row) for row in rd]
        good = {}
        for label, row in entries:
            cid = row["case_id"]
            if current.get(cid) == row["packet_sha256"]:
                if cid in good:
                    raise ValueError(f"{reviewer}: case_id {cid} appears in both {good[cid]} and {label}")
                good[cid] = label
        rows, dropped = [], {}
        for label, row in entries:
            cid = row["case_id"]
            if current.get(cid) == row["packet_sha256"]:
                rows.append(row)
            elif cid in good and round_labels.index(good[cid]) > round_labels.index(label):
                dropped.setdefault(cid, []).append(label)
            else:
                raise ValueError(f"{reviewer}: packet_sha256 for {cid} ({label}) differs from the "
                                 "current render: the pair changed after it was reviewed")
        missing = sorted(set(current) - set(good))
        if missing:
            raise ValueError(f"{reviewer}: rounds {list(round_labels)} do not cover all {len(current)} current "
                             f"pairs; missing {len(missing)}: {missing[:10]}")
        merged[reviewer] = (header, rows)
        superseded[reviewer] = dropped
    if superseded["AG"] != superseded["MO"]:
        raise ValueError(f"the two reviewers' forms supersede different rows: AG {superseded['AG']}, "
                         f"MO {superseded['MO']}")
    out_dir.mkdir(parents=True, exist_ok=True)
    for reviewer, (header, rows) in merged.items():
        with open(out_dir / f"review_{reviewer}.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=header)
            w.writeheader()
            w.writerows(rows)
    (out_dir / MANIFEST).write_text(json.dumps(
        {"pairs_v2_sha256": hashlib.sha256(raw).hexdigest(), "packets": current,
         "merged_rounds": list(round_labels), "superseded": dict(sorted(superseded["AG"].items()))},
        indent=1) + "\n", encoding="utf-8")


def round2_scope(round1_manifest=None):
    """The case_ids round 2 reviews: the pairs unresolved in round 1 plus the round-1 pairs whose current
    packet_sha256 differs from the one the physicians saw. Returns (sorted ids, unresolved ids, changed ids)."""
    if round1_manifest is None:
        round1_manifest = json.loads((VERIFY_DIR / "round1" / MANIFEST).read_text(encoding="utf-8"))
    current = {r["case_id"]: r["packet_sha256"] for r in packet_rows()}
    unresolved = sorted(round1_manifest.get("unresolved") or [])
    changed = sorted(cid for cid, h in round1_manifest["packets"].items() if current.get(cid) != h)
    return sorted(set(unresolved) | set(changed)), unresolved, changed


def _export_round(label, with_author_edits, force=False):
    d = VERIFY_DIR / label
    targets = [d / f"review_{r}.xlsx" for r, _ in REVIEWERS]
    existing = [t.name for t in targets if t.exists()]
    if existing and not force:
        raise RuntimeError(f"refusing to overwrite workbooks that may hold answers: {existing} (use --force)")
    for reviewer, _ in REVIEWERS:
        queue = DATA_V2 / "manual_edit_queue.csv" if (with_author_edits and reviewer == "AG") else None
        export_xlsx(d / f"review_{reviewer}.csv", d / f"review_{reviewer}.xlsx",
                    REVIEWER_NAMES[reviewer], label, queue_csv=queue)


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(prog="r2.verify.build_packets")
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("round1", "round2"):
        p = sub.add_parser(name)
        p.add_argument("--force", action="store_true", help="overwrite existing workbooks")
    p = sub.add_parser("import-edits")
    p.add_argument("xlsx")
    p.add_argument("--author", required=True)
    p.add_argument("--date", required=True)
    p = sub.add_parser("import-review")
    p.add_argument("xlsx")
    p.add_argument("csv")
    p.add_argument("round", choices=("round1", "round2"))
    p = sub.add_parser("merge")
    p.add_argument("rounds", nargs="+")
    a = ap.parse_args(argv)
    if a.cmd == "round1":
        write_forms("round1", allow_unresolved=True)
        _export_round("round1", with_author_edits=True, force=a.force)
    elif a.cmd == "round2":
        write_round_forms("round2", round2_scope()[0])
        _export_round("round2", with_author_edits=False, force=a.force)
    elif a.cmd == "import-edits":
        print(f"{len(import_author_edits(a.xlsx, a.author, a.date))} edits written to data/v2/manual_edits.json")
    elif a.cmd == "import-review":
        import_review_xlsx(a.xlsx, a.csv, a.round)
    elif a.cmd == "merge":
        merge_rounds(a.rounds)


if __name__ == "__main__":
    main()
