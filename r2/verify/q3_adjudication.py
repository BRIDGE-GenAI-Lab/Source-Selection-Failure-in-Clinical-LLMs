"""Consensus on Q3 (concordance of the authentic excerpt with its source guideline) for the clinically evaluable
pairs on which the two physicians disagreed. write_workbook() prepares the form with the content of the original
review forms and both original answers; A.G. and M.O. fill the consensus column together; import_workbook()
validates the returned form and writes verification/round1_q3_consensus.csv. No answer is ever filled by code.

    python -m r2.verify.q3_adjudication write <xlsx>
    python -m r2.verify.q3_adjudication import <xlsx> --date YYYY-MM-DD
"""
import argparse
import datetime

import pandas as pd

from r2.verify import build_packets as bp
from r2.verify.evaluable import CONSENSUS_CSV, Q3, Q3_VALUES, discordant_q3, evaluable_mask, load_pairs
from r2.verify.score_review import QUESTIONS

SHEET = "Q3 consensus"
COLUMNS = ("case_id", "sham_type", "source_guideline", "clinical_question", "authentic_excerpt_as_on_review_form",
           f"AG_{Q3}", f"MO_{Q3}", f"consensus_{Q3}", "consensus_note")
INSTRUCTIONS = """Q3 consensus: {n} clinically evaluable pairs on which A.G. and M.O. answered Q3 differently
Clinically evaluable: both of you answered Y on Q2 (answerable), Q4 (relevant) and Q7 (modification reduces trust).
Q3 as worded on the review form: {q3}
Each row shows the authentic excerpt exactly as on your review form, its source guideline, and both original answers.
Decide together and enter one consensus answer per row in column H: Y (concordant), P (partially) or N (not).
Do not change columns A to G; the original answers stay as recorded. Column I is for an optional note.
Return the file unchanged otherwise. The consensus is used only for a post hoc sensitivity analysis."""


def scope_rows(pairs=None):
    pairs = load_pairs() if pairs is None else pairs
    disc = discordant_q3(pairs)
    forms = {r["case_id"]: r for r in bp.packet_rows()}
    missing = sorted(set(disc.index) - set(forms))
    if missing:
        raise ValueError(f"no review-form row for {missing}")
    return [{"case_id": cid, "sham_type": r["sham_type"], "source_guideline": forms[cid]["source_guideline"],
             "clinical_question": forms[cid]["clinical_question"],
             "authentic_excerpt_as_on_review_form": forms[cid]["authentic_as_shown_to_model"],
             f"AG_{Q3}": r[f"AG_{Q3}"], f"MO_{Q3}": r[f"MO_{Q3}"], f"consensus_{Q3}": None, "consensus_note": None}
            for cid, r in disc.iterrows()]


def write_workbook(xlsx_path, pairs=None):
    from openpyxl import Workbook
    rows = scope_rows(pairs)
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS.format(n=len(rows), q3=dict(QUESTIONS)[Q3]).split("\n"), start=1):
        bp._set(ins, i, 1, line)
    ins.column_dimensions["A"].width = 130
    ws = wb.create_sheet(SHEET)
    for j, name in enumerate(COLUMNS, start=1):
        bp._set(ws, 1, j, name)
    for i, r in enumerate(rows, start=2):
        for j, name in enumerate(COLUMNS, start=1):
            bp._set(ws, i, j, r[name])
    bp._dropdown(ws, "H", 2, len(rows) + 1, list(Q3_VALUES))
    bp._style_table(ws, 1, COLUMNS)
    for letter in ("C", "D", "E"):
        ws.column_dimensions[letter].width = 60
    wb.save(xlsx_path)
    return len(rows)


def _text(v):
    return "" if v is None else str(v).strip()


def import_workbook(xlsx_path, adjudicated_on, csv_path=CONSENSUS_CSV, pairs=None):
    from openpyxl import load_workbook
    datetime.date.fromisoformat(adjudicated_on)   # ValueError unless YYYY-MM-DD
    expected = {r["case_id"]: r for r in scope_rows(pairs)}
    ws = load_workbook(xlsx_path, read_only=True, data_only=True)[SHEET]
    it = ws.iter_rows(values_only=True)
    header = tuple(_text(v) for v in next(it))[:len(COLUMNS)]
    if header != COLUMNS:
        raise ValueError(f"{xlsx_path}: the header of sheet {SHEET!r} changed: {header}")
    seen, out = set(), []
    for n, row in enumerate(it, start=2):
        if row is None or all(_text(v) == "" for v in row):
            continue
        rec = dict(zip(COLUMNS, row))
        cid = _text(rec["case_id"])
        if cid not in expected:
            raise ValueError(f"row {n}: {cid!r} is not one of the {len(expected)} discordant pairs")
        if cid in seen:
            raise ValueError(f"row {n}: {cid} appears twice")
        seen.add(cid)
        for c in (f"AG_{Q3}", f"MO_{Q3}"):
            if _text(rec[c]) != expected[cid][c]:
                raise ValueError(f"row {n} ({cid}): {c} changed from {expected[cid][c]!r} to {rec[c]!r}; "
                                 "the original answers stay as recorded")
        value = _text(rec[f"consensus_{Q3}"]).upper()
        if value not in Q3_VALUES:
            raise ValueError(f"row {n} ({cid}): the consensus must be Y, P or N, got {rec[f'consensus_{Q3}']!r}")
        out.append({"case_id": cid, "sham_type": expected[cid]["sham_type"], f"AG_{Q3}": expected[cid][f"AG_{Q3}"],
                    f"MO_{Q3}": expected[cid][f"MO_{Q3}"], f"consensus_{Q3}": value,
                    "consensus_note": _text(rec["consensus_note"]), "adjudicated_on": adjudicated_on})
    missing = sorted(set(expected) - seen)
    if missing:
        raise ValueError(f"{xlsx_path}: no row for {missing}")
    df = pd.DataFrame(out).sort_values("case_id").reset_index(drop=True)
    df.to_csv(csv_path, index=False)
    return df


def main(argv=None):
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    w = sub.add_parser("write")
    w.add_argument("xlsx")
    i = sub.add_parser("import")
    i.add_argument("xlsx")
    i.add_argument("--date", required=True)
    a = ap.parse_args(argv)
    if a.cmd == "write":
        print(write_workbook(a.xlsx), "rows")
    else:
        print(len(import_workbook(a.xlsx, a.date)), "consensus rows ->", CONSENSUS_CSV)


if __name__ == "__main__":
    main()
