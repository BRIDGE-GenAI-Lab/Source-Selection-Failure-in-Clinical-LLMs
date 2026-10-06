"""Final fix F1: the self-labelling token "fake" in the fabricated-citation edits.

The corrected set must carry no such token (test_build_v2.py); the regenerated round-1 forms must show the
corrected text and leave every other packet unchanged; the audit must disclose the token in the original
model inputs."""
import csv
import hashlib
import json

import pytest

from r2.paths import AUDIT_DIR, DATA_V2, VERIFY_DIR
from r2.stimuli.transforms import self_label_tokens

FC = "fabricated_citation"
ROUND1 = VERIFY_DIR / "round1"
# Pinned from verification/round1/packets_manifest.json at 1a9831f (before F1): digest = sha256 of the
# sorted "case_id:packet_sha256" lines.
HEAD_437_PACKETS = "01b794082bb95f37e79555aad75c41cf29028022c7d8f5032751d43f9d95fc0c"
HEAD_40_FC_PACKETS = "8fa2329f199e314818d00e3cf072e12fb975d7194c92f82d7b9efb6306f73a5c"
PROFILES = {"openai_fc": 4, "gemini_inline7": 1, "deepseek_chat_inline7": 1, "deepseek_reasoner_inline7": 1,
            "together_inline6": 14}


def _digest(pairs):
    return hashlib.sha256("\n".join(f"{a}:{b}" for a, b in sorted(pairs)).encode("utf-8")).hexdigest()


def _pairs():
    return json.loads((DATA_V2 / "pairs_v2.json").read_text(encoding="utf-8"))


def _fc_ids():
    return {c["case_id"] for c in _pairs()["cases"] if c["eval"]["sham_trap_type"] == FC}


# ---------------------------------------------------------------- round-1 forms


def test_round1_manifest_binds_the_corrected_pairs():
    m = json.loads((ROUND1 / "packets_manifest.json").read_text(encoding="utf-8"))
    assert m["pairs_v2_sha256"] == hashlib.sha256((DATA_V2 / "pairs_v2.json").read_bytes()).hexdigest()
    assert len(m["packets"]) == 477 and len(m["unresolved"]) == 23


def test_round1_keeps_437_packets_and_changes_exactly_the_40_fabricated_citation_packets():
    from r2.verify.build_packets import packet_rows
    from r2.stimuli.build_v2 import before_self_label_fix
    packets = json.loads((ROUND1 / "packets_manifest.json").read_text(encoding="utf-8"))["packets"]
    fc = _fc_ids()
    assert len(fc) == 40
    assert _digest([(cid, h) for cid, h in packets.items() if cid not in fc]) == HEAD_437_PACKETS
    # the packets of the 40 pairs as they were before F1, recomputed from the reverted cases
    before = {r["case_id"]: r["packet_sha256"]
              for r in packet_rows([before_self_label_fix(c) for c in _pairs()["cases"] if c["case_id"] in fc])}
    assert _digest(before.items()) == HEAD_40_FC_PACKETS
    assert all(packets[cid] != before[cid] for cid in fc)


@pytest.mark.parametrize("reviewer", ["AG", "MO"])
def test_round1_csv_shows_no_self_label_token(reviewer):
    with open(ROUND1 / f"review_{reviewer}.csv", newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 477
    for r in rows:
        shown = json.loads(r["modified_as_shown_to_model"])
        text = "\n".join(str(x) for v in shown.values() for x in (v if isinstance(v, list) else [v]))
        assert self_label_tokens(text) == [], r["case_id"]
        if r["sham_type"] == FC:
            assert "10.9999/cg.2025." in r["modified_as_shown_to_model"]
            assert "10.9999/fake." not in r["what_changed"]


@pytest.mark.parametrize("reviewer", ["AG", "MO"])
def test_round1_xlsx_matches_the_current_render(reviewer):
    from openpyxl import load_workbook

    from r2.verify.build_packets import REVIEW, _cell_text, _check_review_content
    path = ROUND1 / f"review_{reviewer}.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = [[_cell_text(c) for c in row] for row in wb[REVIEW].iter_rows(values_only=True)]
    while rows and not any(rows[-1]):
        rows.pop()
    _check_review_content(rows, path)  # raises if any content cell differs from a fresh render
    assert len(rows) == 478


def test_round1_files_equal_a_fresh_render(tmp_path, monkeypatch):
    """The committed round-1 CSV, HTML, manifest and README are exactly what build_packets writes for the
    committed pairs (HtmlDiff wraps long lines, so the HTML is compared whole, not searched for a DOI)."""
    import difflib

    from r2.verify import build_packets as bp
    # HtmlDiff numbers its anchors with a class-level counter that every make_table call in the process
    # advances; the committed file was written by a fresh process (python -m r2.verify.build_packets round1)
    monkeypatch.setattr(difflib.HtmlDiff, "_default_prefix", 0)
    raw = (DATA_V2 / "pairs_v2.json").read_bytes()
    d = json.loads(raw)
    bp._write_forms_to(tmp_path, raw, d["cases"], {"unresolved": list(d["unresolved"])})
    for name in ("review_AG.csv", "review_MO.csv", "packets.html", "packets_manifest.json", "README_REVIEWERS.md"):
        assert (tmp_path / name).read_bytes() == (ROUND1 / name).read_bytes(), name
    assert "fake" not in (ROUND1 / "packets.html").read_text(encoding="utf-8").lower()


# ---------------------------------------------------------------- audit disclosure


@pytest.fixture(scope="module")
def committed():
    return json.loads((AUDIT_DIR / "self_label_tokens.json").read_text(encoding="utf-8"))


def test_audit_counts_840_fake_rows_all_fabricated_citation(committed):
    fake = committed["modified_tool_payload"]["fake"]
    assert fake["rows"] == 840 and fake["n_cases"] == 40 and fake["n_models"] == 21
    assert fake["by_sham_type"] == {FC: 840}
    assert fake["by_profile"] == {p: 40 * n for p, n in PROFILES.items()}
    assert sorted(fake["case_ids"]) == sorted(_fc_ids())
    assert committed["n_rows"] == 10500 and committed["n_models"] == 21
    for tok in committed["tokens"]:
        assert committed["authentic_tool_payload"][tok]["rows"] == 0, tok
        if tok != "fake":
            assert committed["modified_tool_payload"][tok]["rows"] == 0, tok
            assert committed["prompt_text"][tok]["rows"] == 0, tok


def test_audit_discloses_the_openai_instruction_mention(committed):
    p = committed["prompt_text"]["fake"]
    assert p["rows"] == 2000 and p["by_profile"] == {"openai_fc": 2000}
    assert sum(p["by_sham_type"].values()) == 2000
    assert any("fake citations" in c for c in p["contexts"])


def test_audit_json_is_current(committed):
    from r2.audit.self_label_tokens import compute
    assert compute() == committed


def test_audit_readme_has_the_section():
    text = (AUDIT_DIR / "README.md").read_text(encoding="utf-8")
    assert "## self_label_tokens.json" in text
    assert "840" in text and "cue-confounded" in text and "10.9999/fake." in text and "10.9999/cg." in text
