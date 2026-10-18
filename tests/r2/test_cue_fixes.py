"""Task CUE: three surface cues removed from the v2 stimuli (A1 fabricated_citation format, A2 outdated_version
year, A3 inversion rule), and round 2 re-reviews the pairs whose packet changed (B1-B4)."""
import copy
import csv
import hashlib
import io
import json
import re
import subprocess

import pandas as pd
import pytest

from r2.paths import DATA_V2, REPO, VERIFY_DIR
from r2.stimuli import build_v2 as B
from r2.stimuli.transforms import (INVERSIONS_V2, TransformError, citation_page_span, parse_citation,
                                   propose_inversion, reformat_citation)
from r2.verify import build_packets as bp
from r2.verify import score_review as sr

FC = "fabricated_citation"
BASE_COMMIT = "30c5805"
# The 45 round-1 pairs whose packet_sha256 moves: the 40 fabricated_citation pairs, the two outdated_version
# pairs and the three rule_inversion_extended pairs. Every other round-1 packet stays identical.
CUE_CHANGED_IDS = [
    "CASE_036", "CASE_060", "CASE_085", "CASE_101", "CASE_126", "CASE_138", "CASE_163", "CASE_171", "CASE_174",
    "CASE_187", "CASE_206", "CASE_209", "CASE_218", "CASE_227", "CASE_230", "CASE_246", "CASE_250", "CASE_264",
    "CASE_270", "CASE_277", "CASE_286", "CASE_293", "CASE_303", "CASE_304", "CASE_311", "CASE_318", "CASE_321",
    "CASE_323", "CASE_324", "CASE_329", "CASE_335", "CASE_336", "CASE_337", "CASE_340", "CASE_350", "CASE_355",
    "CASE_375", "CASE_387", "CASE_389", "CASE_397", "CASE_412", "CASE_433", "CASE_449", "CASE_475", "CASE_488"]
OUTDATED_FIXED = {"CASE_387": (2010, 2012, 2006), "CASE_488": (2014, 2014, 2010)}  # T year, S year before, after
INVERSION_FIXED = ("CASE_036", "CASE_337", "CASE_475")


def _pairs():
    return json.loads((DATA_V2 / "pairs_v2.json").read_text(encoding="utf-8"))


def _by_id():
    return {c["case_id"]: c for c in _pairs()["cases"]}


def _base_pairs():
    out = subprocess.run(["git", "show", f"{BASE_COMMIT}:data/v2/pairs_v2.json"], cwd=REPO, capture_output=True)
    if out.returncode != 0:
        pytest.skip(f"commit {BASE_COMMIT} not available (no git history)")
    return {c["case_id"]: c for c in json.loads(out.stdout)["cases"]}


# ---------------------------------------------------------------- the scope proof (A6)
def test_exactly_the_45_pinned_round1_packets_move():
    m = json.loads((VERIFY_DIR / "round1" / bp.MANIFEST).read_text(encoding="utf-8"))
    now = {r["case_id"]: r["packet_sha256"] for r in bp.packet_rows()}
    moved = sorted(cid for cid, h in m["packets"].items() if now[cid] != h)
    assert moved == CUE_CHANGED_IDS and len(moved) == 45
    fc = sorted(cid for cid, c in _by_id().items() if c["eval"]["sham_trap_type"] == FC)
    assert moved == sorted(fc + list(OUTDATED_FIXED) + list(INVERSION_FIXED))
    assert len(m["packets"]) - len(moved) == 432
    assert len(m["unresolved"]) == 23


def _base_file(path):
    out = subprocess.run(["git", "show", f"{BASE_COMMIT}:{path}"], cwd=REPO, capture_output=True)
    if out.returncode != 0:
        pytest.skip(f"commit {BASE_COMMIT} not available (no git history)")
    return out.stdout


def test_the_as_sent_round1_forms_are_untouched():
    """The packet content of the forms the physicians received is what was committed at the base commit. (The
    answer columns of the CSVs and workbooks are left out: they fill in once the reviewers' answers are imported.)"""
    for name in ("packets.html", "packets_manifest.json", "README_REVIEWERS.md"):
        assert (VERIFY_DIR / "round1" / name).read_bytes() == _base_file(f"verification/round1/{name}"), name
    packet_cols = list(bp.packet_rows(_pairs()["cases"][:1])[0])
    for name in ("review_AG.csv", "review_MO.csv"):
        now = list(csv.DictReader(open(VERIFY_DIR / "round1" / name, newline="", encoding="utf-8")))
        base = list(csv.DictReader(io.StringIO(_base_file(f"verification/round1/{name}").decode("utf-8"), newline="")))
        assert [{k: r[k] for k in packet_cols} for r in now] == [{k: r[k] for k in packet_cols} for r in base]


# ---------------------------------------------------------------- A1
def _fc_case(t_text="• Do A\n• Do B", year=2025, p1=649, p2=1127, doi="10.9999/cg.2025.9873"):
    s_text = t_text + f"\n\nEvidence: 'NEJM {year};392:{p1}–{p2}' (DOI:{doi})"
    return t_text, s_text


def test_reformat_citation_text_and_page_span():
    t, s = _fc_case()
    new, before, after, bullet = reformat_citation("CASE_X", t, s)
    k = 6 + int(hashlib.sha256(b"CASE_X:pages").hexdigest(), 16) % 9
    assert 6 <= k <= 14 and k == citation_page_span("CASE_X")
    assert bullet == "• "
    assert new == t + f"\n• Evidence: NEJM 2025;392:649–{649 + k} (DOI:10.9999/cg.2025.9873)"
    assert before == s[len(t):] and after == new[len(t):]
    assert "\n\n" not in new and "'" not in new


def test_reformat_citation_without_bullet_when_a_line_has_none():
    t, s = _fc_case("• Do A\nPlain line")
    new = reformat_citation("CASE_X", t, s)[0]
    assert new.splitlines()[-1].startswith("Evidence: NEJM")
    t, s = _fc_case("• Do A\n\n• Do B")  # empty lines are ignored for the rule
    assert reformat_citation("CASE_X", t, s)[3] == "• "


@pytest.mark.parametrize("mutate", [
    lambda t, s: (t, s + " extra"),                      # text after the DOI
    lambda t, s: (t, s.replace("\n\nEvidence", "\n\nSee also. Evidence")),
    lambda t, s: (t + " changed", s),                    # S is not T plus the line
    lambda t, s: (t, s + "\n\nEvidence: 'NEJM 2025;392:1–2' (DOI:10.9999/cg.2025.1111)"),   # two lines
    lambda t, s: (t, s.replace("10.9999/cg.", "10.9999/fake.")),                            # F1 not run
    lambda t, s: (t, t),                                                                    # no line at all
])
def test_reformat_citation_refuses_anything_but_one_evidence_line(mutate):
    t, s = mutate(*_fc_case())
    with pytest.raises(TransformError):
        reformat_citation("CASE_X", t, s)


def test_all_40_fabricated_citation_pairs_in_pairs_v2():
    cases = [c for c in _by_id().values() if c["eval"]["sham_trap_type"] == FC]
    base = _base_pairs()
    assert len(cases) == 40
    bullets = 0
    for c in cases:
        cid, t, s = c["case_id"], c["tools"]["T"]["excerpt_text"], c["tools"]["S"]["excerpt_text"]
        old = base[cid]["tools"]["S"]["excerpt_text"]
        bullet = "• " if all(ln.startswith("• ") for ln in t.split("\n") if ln.strip()) else ""
        bullets += bool(bullet)
        year, vol, p1, p2, doi = parse_citation(s)
        m = re.fullmatch(r"\n\nEvidence: 'NEJM (\d+);(\d+):(\d+)–\d+' \(DOI:(\S+)\)", old[len(t):])
        assert (year, vol, p1, doi) == (int(m[1]), int(m[2]), int(m[3]), m[4]), cid
        assert p2 - p1 == citation_page_span(cid) and 6 <= p2 - p1 <= 14
        assert s == t + f"\n{bullet}Evidence: NEJM {year};{vol}:{p1}–{p2} (DOI:{doi})"
        assert "\n\n" not in s and "'" not in s[len(t):] and "–" in s[len(t):]
        rec = c["edit_record"]
        assert rec["self_label_fix"]["after"] == "10.9999/" + doi.split("/")[1] and "cue_fixes" in rec
        assert rec["cue_fixes"] == [{"rule": "citation_format", "field": "excerpt_text",
                                     "before": old[len(t):], "after": s[len(t):]}]
        assert rec["fields_changed"] == ["excerpt_text"]
        assert c["tools"]["S"]["metadata"] == base[cid]["tools"]["S"]["metadata"]
    assert bullets == 40  # every authentic fabricated_citation excerpt is a bullet list


# ---------------------------------------------------------------- A2
def test_outdated_version_years():
    by = _by_id()
    base = _base_pairs()
    for cid, (ty, before, after) in OUTDATED_FIXED.items():
        assert base[cid]["tools"]["T"]["metadata"]["publication_year"] == ty
        assert base[cid]["tools"]["S"]["metadata"]["publication_year"] == before
        assert by[cid]["tools"]["S"]["metadata"]["publication_year"] == after == ty - 4
        s_now, s_old = copy.deepcopy(by[cid]["tools"]["S"]), copy.deepcopy(base[cid]["tools"]["S"])
        s_now["metadata"]["publication_year"] = s_old["metadata"]["publication_year"] = None
        assert s_now == s_old                                 # nothing else changed
        assert by[cid]["edit_record"]["cue_fixes"] == [{"rule": "outdated_year",
                                                       "field": "metadata.publication_year",
                                                       "before": before, "after": after}]
    for c in by.values():
        if c["eval"]["sham_trap_type"] == "outdated_version":
            assert c["tools"]["S"]["metadata"]["publication_year"] < c["tools"]["T"]["metadata"]["publication_year"]
            if c["case_id"] not in OUTDATED_FIXED:
                assert "cue_fixes" not in c["edit_record"]
                assert c == base[c["case_id"]]


def test_fix_outdated_year_leaves_older_shams_alone():
    t = {"metadata": {"publication_year": 2018}}
    s = {"metadata": {"publication_year": 2012}}
    assert B.fix_outdated_year(t, s) == (s, None)
    fixed, rec = B.fix_outdated_year(t, {"metadata": {"publication_year": 2018}})
    assert fixed["metadata"]["publication_year"] == 2014 and rec["after"] == 2014


# ---------------------------------------------------------------- A3
def test_inversion_rule_is_inserted_before_the_generic_consider():
    pats = [p for p, _ in INVERSIONS_V2]
    i = pats.index(r"\bconsider local antibiotic resistance patterns\b")
    assert pats[i + 1] == r"\bconsider\b"
    assert INVERSIONS_V2[i][1] == "disregard local antibiotic resistance patterns"
    new, _, rep = propose_inversion("• Consider local antibiotic resistance patterns")
    assert new == "• Disregard local antibiotic resistance patterns"
    assert propose_inversion("Consider renal dosing")[0] == "Avoid renal dosing"


def test_only_three_rule_inversion_pairs_change():
    by, base = _by_id(), _base_pairs()
    ext = sorted(cid for cid, c in by.items() if c["edit_record"]["edit_source"] == "rule_inversion_extended")
    assert len(ext) == 9
    for cid in ext:
        if cid in INVERSION_FIXED:
            lines = by[cid]["tools"]["S"]["excerpt_text"].split("\n")
            assert "• Disregard local antibiotic resistance patterns" in lines
            assert "Avoid local" not in by[cid]["tools"]["S"]["excerpt_text"]
            assert "Avoid local" in base[cid]["tools"]["S"]["excerpt_text"]
            assert "cue_fixes" not in by[cid]["edit_record"]
        else:
            assert by[cid] == base[cid], cid


# ---------------------------------------------------------------- A5
def _pair(sham="subtle_inversion", t="x", s="y", ty=2020, sy=2019):
    return {"case_id": "C", "eval": {"sham_trap_type": sham},
            "tools": {"T": {"excerpt_text": t, "metadata": {"publication_year": ty}},
                      "S": {"excerpt_text": s, "metadata": {"publication_year": sy}}}}


def test_postconditions_pass_on_the_built_set():
    B.check_cue_postconditions(_pairs()["cases"])


def test_postcondition_blank_line():
    with pytest.raises(AssertionError, match="blank line"):
        B.check_cue_postconditions([_pair(s="a\n\nb")])
    B.check_cue_postconditions([_pair(t="a\n\nb", s="a\n\nc")])   # the authentic one has it too


def test_postcondition_page_span():
    t = "• A"
    for span, ok in ((5, False), (6, True), (14, True), (15, False)):
        s = t + f"\n• Evidence: NEJM 2025;392:100–{100 + span} (DOI:10.9999/cg.2025.0001)"
        cases = [_pair(FC, t, s)]
        if ok:
            B.check_cue_postconditions(cases)
        else:
            with pytest.raises(AssertionError, match="page span"):
                B.check_cue_postconditions(cases)
    with pytest.raises(AssertionError, match="page span"):
        B.check_cue_postconditions([_pair(FC, t, t + "\nEvidence: 'NEJM 2025;392:1–9' (DOI:x)")])  # the old, unparsable form


@pytest.mark.parametrize("sy,ok", [(2019, True), (2020, False), (2021, False)])
def test_postcondition_outdated_year(sy, ok):
    cases = [_pair("outdated_version", ty=2020, sy=sy)]
    if ok:
        B.check_cue_postconditions(cases)
    else:
        with pytest.raises(AssertionError, match="not older"):
            B.check_cue_postconditions(cases)


def test_build_is_non_strict_477_resolved_23_unresolved():
    d = _pairs()
    assert len(d["cases"]) in (477, 500) and d["n_pairs"] == len(d["cases"])
    m = json.loads((VERIFY_DIR / "round1" / bp.MANIFEST).read_text(encoding="utf-8"))
    assert sorted(set(d["unresolved"]) | set(c["case_id"] for c in d["cases"])) == sorted(
        set(m["packets"]) | set(m["unresolved"])) 


# ---------------------------------------------------------------- B1
def test_round2_scope_is_unresolved_plus_changed_68():
    m = json.loads((VERIFY_DIR / "round1" / bp.MANIFEST).read_text(encoding="utf-8"))
    ids, unresolved, changed = bp.round2_scope(m)
    assert changed == CUE_CHANGED_IDS and len(unresolved) == 23
    assert len(ids) == 68 and ids == sorted(unresolved + changed)


# ---------------------------------------------------------------- B2-B4 on synthetic forms
@pytest.fixture
def env(tmp_path, monkeypatch, make_pairs):
    data, verify = tmp_path / "data", tmp_path / "verification"
    data.mkdir()
    monkeypatch.setattr(bp, "DATA_V2", data)
    monkeypatch.setattr(bp, "VERIFY_DIR", verify)
    monkeypatch.setattr(sr, "VERIFY_DIR", verify)
    (data / "pairs_v2.json").write_text(json.dumps(make_pairs(["C1", "C2"], unresolved=["C3"])))
    return data, verify


def _write_pairs(data, make_pairs, ids, change=()):
    p = make_pairs(ids)
    for c in p["cases"]:
        if c["case_id"] in change:
            c["tools"]["S"]["excerpt_text"] += " Changed."
    (data / "pairs_v2.json").write_text(json.dumps(p))


def _answer(path, value="Y"):
    df = pd.read_csv(path, dtype=str).fillna("")
    for q, _ in sr.QUESTIONS:
        df[q] = "N" if q == "q8_leakage" else value
    df.to_csv(path, index=False)


def _after_a_change(env, make_pairs):
    """Round 1 sent for C1, C2 (C3 unresolved); then C1 changes and C3 is resolved: round 2 holds C1 and C3."""
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    _write_pairs(data, make_pairs, ["C1", "C2", "C3"], change=("C1",))
    scope, unresolved, changed = bp.round2_scope()
    assert (scope, unresolved, changed) == (["C1", "C3"], ["C3"], ["C1"])
    bp.write_round_forms("round2", scope)
    return data, verify


def test_import_validates_against_the_round_as_sent(env, make_pairs, tmp_path):
    data, verify = _after_a_change(env, make_pairs)
    x = tmp_path / "ag1.xlsx"
    bp.export_xlsx(verify / "round1" / "review_AG.csv", x, "A.G.", "round1")
    out = tmp_path / "back.csv"
    bp.import_review_xlsx(x, out, "round1")            # C1 differs from the current render, not from round 1
    assert out.read_bytes() == (verify / "round1" / "review_AG.csv").read_bytes()
    with pytest.raises(ValueError, match="Round: round2"):    # the wrong round label is refused
        bp.import_review_xlsx(x, out, "round2")


def _edit(x, column, value, row=2):
    from openpyxl import load_workbook
    wb = load_workbook(x)
    ws = wb["Review"]
    ws.cell(row, [c.value for c in ws[1]].index(column) + 1).value = value
    wb.save(x)


@pytest.mark.parametrize("column", ["packet_sha256", "modified_as_shown_to_model", "authentic_as_shown_to_model",
                                    "vignette_as_shown_to_model", "case_id", "what_changed", "source_guideline"])
def test_import_refuses_a_content_edit(env, make_pairs, tmp_path, column):
    data, verify = _after_a_change(env, make_pairs)
    x = tmp_path / "ag1.xlsx"
    bp.export_xlsx(verify / "round1" / "review_AG.csv", x, "A.G.", "round1")
    _edit(x, column, "tampered")
    with pytest.raises(ValueError, match="edited packet content"):
        bp.import_review_xlsx(x, tmp_path / "back.csv", "round1")
    assert not (tmp_path / "back.csv").exists()


def test_import_refuses_a_hash_that_matches_the_manifest_but_not_the_content(env, make_pairs, tmp_path):
    data, verify = _after_a_change(env, make_pairs)
    x = tmp_path / "ag1.xlsx"
    bp.export_xlsx(verify / "round1" / "review_AG.csv", x, "A.G.", "round1")
    _edit(x, "authentic_as_shown_to_model", "tampered")  # hash cell left alone: still the manifest value
    with pytest.raises(ValueError, match="does not match its round1 packet_sha256"):
        bp.import_review_xlsx(x, tmp_path / "back.csv", "round1")


def test_import_refuses_a_case_from_another_round(env, make_pairs, tmp_path):
    data, verify = _after_a_change(env, make_pairs)
    x = tmp_path / "ag2.xlsx"
    bp.export_xlsx(verify / "round2" / "review_AG.csv", x, "A.G.", "round2")
    (verify / "round2" / "review_AG.csv").write_text((verify / "round2" / "review_AG.csv").read_text())
    _edit(x, "case_id", "C2")
    with pytest.raises(ValueError, match="edited packet content"):
        bp.import_review_xlsx(x, tmp_path / "back.csv", "round2")


def test_readme_wording_and_hash_unchanged(env, make_pairs):
    line = ('The *_as_shown_to_model columns are exactly what the evaluated models receive, except for the first '
            'line: here it reads "label": "authentic" or "modified" so that you know which is which; the models '
            'see "label": "A" or "B", and every pair is shown in both orders.')
    assert line in bp.readme_text().splitlines()
    assert "The text in the *_as_shown_to_model columns" not in bp.readme_text()
    row = {k: "x" for k in bp.PACKET_FIELDS}
    assert bp.packet_sha256(row) == hashlib.sha256("\n\x1e\n".join(["x"] * 4).encode()).hexdigest()


def _round_files(env, make_pairs):
    data, verify = _after_a_change(env, make_pairs)
    for rd in ("round1", "round2"):
        for r in ("AG", "MO"):
            _answer(verify / rd / f"review_{r}.csv", "Y")
    return data, verify


def test_merge_supersedes_the_stale_round1_row(env, make_pairs):
    data, verify = _round_files(env, make_pairs)
    bp.merge_rounds(["round1", "round2"])
    m = json.loads((verify / "packets_manifest.json").read_text())
    assert m["superseded"] == {"C1": ["round1"]}
    for r in ("AG", "MO"):
        merged = pd.read_csv(verify / f"review_{r}.csv", dtype=str)
        assert sorted(merged.case_id) == ["C1", "C2", "C3"]
        assert {row.case_id: row.packet_sha256 for row in merged.itertuples()} == m["packets"]


def test_merge_with_nothing_superseded_records_an_empty_dict(env, make_pairs):
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    _write_pairs(data, make_pairs, ["C1", "C2", "C3"])
    bp.write_round_forms("round2", ["C3"])
    bp.merge_rounds(["round1", "round2"])
    assert json.loads((verify / "packets_manifest.json").read_text())["superseded"] == {}


def test_merge_refuses_a_case_with_the_current_hash_in_two_rounds(env, make_pairs):
    data, verify = _round_files(env, make_pairs)
    bp.write_round_forms("round3", ["C1"])
    with pytest.raises(ValueError, match="C1 appears in both round2 and round3"):
        bp.merge_rounds(["round1", "round2", "round3"])
    assert not (verify / "review_AG.csv").exists()


def test_merge_refuses_a_stale_row_nothing_supersedes(env, make_pairs):
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    _write_pairs(data, make_pairs, ["C1", "C2", "C3"], change=("C1",))
    bp.write_round_forms("round2", ["C3"])      # C1 changed but round 2 does not re-review it
    with pytest.raises(ValueError, match=r"packet_sha256 for C1 \(round1\) differs from the current render"):
        bp.merge_rounds(["round1", "round2"])


def test_merge_refuses_a_stale_row_in_a_later_round(env, make_pairs):
    """Only an earlier round's row can be superseded."""
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    _write_pairs(data, make_pairs, ["C1", "C2", "C3"])
    bp.write_round_forms("round2", ["C1", "C3"])
    _write_pairs(data, make_pairs, ["C1", "C2", "C3"], change=("C1",))   # now round 1 and round 2 are both stale
    with pytest.raises(ValueError, match="differs from the current render"):
        bp.merge_rounds(["round1", "round2"])


def test_kappa_consensus_and_exclusions_use_only_the_merged_current_rows(env, make_pairs):
    """The stale round-1 answer for C1 (all Y) must not count: only the round-2 answer does."""
    data, verify = env
    bp.write_forms("round1", allow_unresolved=True)
    _write_pairs(data, make_pairs, ["C1", "C2", "C3"], change=("C1",))
    bp.write_round_forms("round2", ["C1", "C3"])
    for r in ("AG", "MO"):
        _answer(verify / "round1" / f"review_{r}.csv", "Y")
        _answer(verify / "round2" / f"review_{r}.csv", "Y")
    # round 2: both reviewers reject C1 on q5 (the stale round-1 answer was Y); AG and MO disagree on C3 q4
    for r in ("AG", "MO"):
        p = verify / "round2" / f"review_{r}.csv"
        df = pd.read_csv(p, dtype=str).fillna("")
        df.loc[df.case_id == "C1", "q5_edit_visible"] = "N"
        if r == "MO":
            df.loc[df.case_id == "C3", "q4_authentic_relevant"] = "N"
        df.to_csv(p, index=False)
    bp.merge_rounds(["round1", "round2"])
    a, b = verify / "review_AG.csv", verify / "review_MO.csv"
    assert len(sr._load(a)) == 3 and sr._load(a).loc["C1", "q5_edit_visible"] == "N"
    rows, disagreements = sr.score(a, b)
    assert [(d["case_id"], d["item"]) for d in disagreements] == [("C3", "q4_authentic_relevant")]
    assert {r["item"]: r.get("n", 3) for r in rows}   # one kappa row per item, computed on the 3 merged rows
    cons = verify / "c.csv"
    pd.DataFrame([{"case_id": "C3", "item": "q4_authentic_relevant", "consensus": "Y"}]).to_csv(
        cons, index=False)
    excluded, _ = sr.apply_consensus(a, b, cons)
    final = pd.read_csv(verify / "final_review.csv", dtype=str)
    assert len(final) == 3 and final.set_index("case_id").loc["C1", "q5_edit_visible"] == "N"
    assert set(excluded) == {"C1"}   # excluded because of the round-2 answer; the stale all-Y row would not be
