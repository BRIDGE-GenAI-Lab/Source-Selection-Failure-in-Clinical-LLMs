"""End-to-end rehearsal of r2.pipeline in a fresh sandbox, with simulated physicians and a simulated
OpenAI-compatible API. NOTHING IT PRODUCES IS DATA.

    python -m r2.rehearse <new sandbox dir> [--models N] [--no-render] [--seed S]

The parent process (run without R2_SANDBOX) refuses any directory that r2.analysis.synthetic.assert_sandbox
refuses, hashes the real state (data/v2, verification, analysis_repeated and r2_runs when present, the
npj_R2_WORKING folder, the git HEAD, tags and status of the repository), creates the sandbox with
r2.pipeline.sandbox_init, runs every step in ONE child process with R2_SANDBOX=<dir> and R2_REPO=<this
checkout> (so every r2.paths write location is inside the sandbox), hashes the real state again and fails if
anything changed. Steps of the child:
  1. simulated author edits for the queued pairs (author "SIMULATED"), written into a copy of the round-1 AG
     workbook, then pipeline import-edits;
  2. simulated physician answers for round 1 and round 2 (seeded; mostly Y, with N and P answers so that there
     are exclusions, flags and disagreements), pipeline consensus, the simulated consensus file for every
     disagreement, pipeline consensus again;
  3. a simulated pilot change (twice the max_tokens of one non-reasoning model, so the documents report a range);
     a fake OpenAI-compatible client injected through the harness's client factory: probe, pilot (10 pairs),
     and after the freeze the main run (every model x 500 pairs x 2 orders) and its resume. The client has a
     position bias per model, invalid replies, <think> tags, 400 and 429 errors with Retry-After headers, and
     one cell whose 400 error persists (it ends as persistent_error);
  4. pipeline pilot-summary; pipeline freeze --dry-run, then the freeze itself in the sandbox freeze repository,
     and the simulated public record of the push and of the release;
  5. pipeline post-results, then pipeline package-runs (the run-log manifest and tarball);
  6. rehearsal_report.md in the sandbox: each step's result, the remaining markers by class, every NUMERIC
     marker not filled, the fake API events and the timing.
"""
import argparse
import csv
import hashlib
import json
import os
import subprocess
import sys
import time
from collections import Counter
from datetime import datetime, timedelta, timezone
from email.utils import format_datetime
from pathlib import Path
from types import SimpleNamespace

SEED = 20261005
SIM_AUTHOR = "SIMULATED"
SIM_TAG_URL = "https://github.com/example/r2-rehearsal-SIMULATED/releases/tag/r2-analysis-plan-frozen"
SIM_RELEASE = "r2-rehearsal-SIMULATED"
CANONICAL_REPO = Path.home() / "npj_r2" / "r2_repo"
REPORT = "rehearsal_report.md"
CHILD_RESULT = "rehearsal_result.json"


def _u(*parts):
    """A deterministic number in [0, 1) from the parts."""
    h = hashlib.sha256("\x1f".join(map(str, parts)).encode("utf-8")).digest()
    return int.from_bytes(h[:8], "big") / 2 ** 64


# ================================================================ real-state guard (parent)
def _tree_hashes(root, only_suffix=None):
    root = Path(root)
    out = {}
    if not root.exists():
        return out
    for p in sorted(root.rglob("*")):
        if not p.is_file() or p.name.startswith("._") or p.name == ".DS_Store" or "__pycache__" in p.parts:
            continue
        if only_suffix and p.suffix != only_suffix:
            continue
        out[str(p)] = hashlib.sha256(p.read_bytes()).hexdigest()
    return out


def _git_state(repo):
    """HEAD, tags and status of `repo`, read with no GIT_* variables (so the guard reads this checkout, not a
    repository an inherited GIT_DIR names) and no optional locks."""
    from r2.sandbox_git import clean_env

    def g(*a):
        return subprocess.run(["git", *a], cwd=repo, capture_output=True, text=True, env=clean_env()).stdout
    return {"head": g("rev-parse", "HEAD").strip(), "tags": g("tag", "--list").split(),
            "status": g("status", "--porcelain").splitlines()}


def real_state():
    """Hashes of every real location the pipeline could write: data/v2, verification, analysis_repeated and
    r2_runs (when present) of this checkout and of the canonical checkout, the whole npj_R2_WORKING folder
    (its .docx files included), and the git state of both checkouts."""
    from r2 import paths
    repos = [Path(paths.REPO)]
    if CANONICAL_REPO.exists() and CANONICAL_REPO.resolve() != Path(paths.REPO).resolve():
        repos.append(CANONICAL_REPO)
    state = {"files": {}, "git": {}, "roots": []}
    for repo in repos:
        for rel in ("data/v2", "verification", "analysis_repeated"):
            state["roots"].append(str(repo / rel))
            state["files"].update(_tree_hashes(repo / rel))
        state["git"][str(repo)] = _git_state(repo)
    for root in (paths.REAL_RUNS_V2, paths.REAL_RUNS_ARCHIVE, paths.REAL_R2_WORKING):
        state["roots"].append(str(root))
        state["files"].update(_tree_hashes(root))
    state["docx"] = {k: v for k, v in state["files"].items() if k.endswith(".docx")}
    return state


def compare_state(before, after):
    files_b, files_a = before["files"], after["files"]
    changed = sorted(k for k in set(files_b) & set(files_a) if files_b[k] != files_a[k])
    return {"unchanged": not changed and set(files_b) == set(files_a) and before["git"] == after["git"],
            "n_files": len(files_b), "n_docx": len(before["docx"]), "changed": changed,
            "added": sorted(set(files_a) - set(files_b)), "removed": sorted(set(files_b) - set(files_a)),
            "git_changed": [r for r in before["git"] if before["git"][r] != after["git"].get(r)],
            "roots": before["roots"]}


# ================================================================ simulated physicians
class Physicians:
    """Seeded answers: a true answer per pair and item (mostly Y; some N, P and leakage Y), and each reviewer
    reports it with a 2.5% chance of another allowed answer. The consensus is the true answer."""

    def __init__(self, seed=SEED):
        self.seed = seed

    def truth(self, cid, item):
        from r2.verify.score_review import EXCLUDE_IF_NO
        u = _u(self.seed, "truth", cid, item)
        if item == "q8_leakage":
            return "Y" if u < 0.012 else "N"
        if item == "q3_authentic_concordant" and u > 0.98:
            return "P"
        if item in EXCLUDE_IF_NO:
            return "N" if u < 0.012 else "Y"
        return "N" if u < 0.01 else "Y"

    def answer(self, reviewer, cid, item):
        from r2.verify.score_review import allowed_answers
        t = self.truth(cid, item)
        if _u(self.seed, "slip", reviewer, cid, item) < 0.025:
            others = [a for a in allowed_answers(item) if a != t]
            return others[int(_u(self.seed, "pick", reviewer, cid, item) * len(others))]
        return t

    def note(self, reviewer, cid, answers):
        bad = [q for q, a in answers.items() if (a == "N") != (q == "q8_leakage") or a == "P"]
        return f"SIMULATED note ({reviewer}): {', '.join(bad)}" if bad else ""


def fill_review_workbook(src, dst, reviewer, phys, author_edits=None):
    """A copy of a review workbook with every answer cell filled (and, for A.G. in round 1, the Author_edits
    sheet). Only answer columns and the find/replace/field cells change."""
    from openpyxl import load_workbook
    from r2.verify import build_packets as BP
    from r2.verify.score_review import QUESTIONS
    wb = load_workbook(src)
    ws = wb[BP.REVIEW]
    header = [c.value for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(header)}
    n = 0
    for r in range(2, ws.max_row + 1):
        cid = ws.cell(row=r, column=col["case_id"]).value
        if not cid:
            continue
        answers = {q: phys.answer(reviewer, cid, q) for q, _ in QUESTIONS}
        for q, a in answers.items():
            BP._set(ws, r, col[q], a)
        note = phys.note(reviewer, cid, answers)
        if note:
            BP._set(ws, r, col["notes"], note)
        n += 1
    if author_edits is not None:
        ae = wb[BP.AUTHOR_EDITS]
        head = [c.value for c in ae[2]]
        acol = {name: i + 1 for i, name in enumerate(head)}
        for r in range(3, ae.max_row + 1):
            cid = ae.cell(row=r, column=acol["case_id"]).value
            if cid in author_edits:
                e = author_edits[cid]
                for k in ("field", "find", "replace"):
                    BP._set(ae, r, acol[k], e[k])
                BP._set(ae, r, acol["rationale"], "SIMULATED edit for the rehearsal; not a physician's edit")
    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return n


VERBS = ("Apply", "Assess", "Check", "Immobilize", "Refer", "Use", "Treat", "Obtain", "Identify", "Continue",
         "Target", "Give", "Start", "Administer", "Consider", "Monitor", "Direct")


def author_edits_for(queue, manual_path):
    """The edits the simulated A.G. writes into the Author_edits sheet. When `manual_path` (the sandbox's
    data/v2/manual_edits.json) exists, the real physician edits are already imported, and exactly the stored
    field/find/replace are written, so import-edits takes its idempotent path; otherwise synthetic ones."""
    manual_path = Path(manual_path)
    if not manual_path.exists():
        return simulated_edits(queue)
    stored = json.loads(manual_path.read_text(encoding="utf-8"))
    missing = sorted(set(queue) - set(stored))
    if missing:
        raise RuntimeError(f"{manual_path} has no edit for queued pairs {missing}")
    return {cid: {k: stored[cid][k] for k in ("field", "find", "replace")} for cid in queue}


def relabel_stored_edits(manual_path, author=None):
    """Sandbox only: the copied physician edits are rehearsal data now, so they carry the simulated author, as
    the documents' sandbox mode expects (edits/fill.py EXPECTED_EDIT_AUTHOR). Returns how many were relabelled."""
    from r2 import paths
    if paths.SANDBOX is None:
        raise RuntimeError("relabel_stored_edits only runs inside a sandbox")
    manual_path = Path(manual_path)
    if not manual_path.exists():
        return 0
    stored = json.loads(manual_path.read_text(encoding="utf-8"))
    for e in stored.values():
        e["author"] = author or SIM_AUTHOR
    manual_path.write_text(json.dumps(stored, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    return len(stored)


def simulated_edits(queue):
    """One valid edit of the declared type per queued pair: subtle_inversion negates one bullet of the excerpt
    (a bullet that occurs once), authority_mimicry writes a look-alike organization name."""
    from r2.io import load_original_cases
    originals, _ = load_original_cases()
    out = {}
    for cid in queue:
        case = originals[cid]
        sham = case["eval"]["sham_trap_type"]
        s = case["tools"]["S"]
        if sham == "authority_mimicry":
            org = s["metadata"]["organization"]
            out[cid] = {"field": "metadata.organization", "find": org,
                        "replace": org.replace(" ", "-", 1) if " " in org else org + "-"}
            continue
        text = s["excerpt_text"]
        for line in text.split("\n"):
            body = line.lstrip("• ").strip()
            if len(body) < 12 or text.count(body) != 1:
                continue
            first = body.split()[0]
            new = ("Do not " + first.lower() + body[len(first):]) if first in VERBS else body + " is not recommended"
            out[cid] = {"field": "excerpt_text", "find": body, "replace": new}
            break
        else:
            raise RuntimeError(f"{cid}: no bullet to invert")
    return out


# ================================================================ simulated API
class _HTTPResponse:
    def __init__(self, status, headers=None):
        self.status_code, self.headers = status, dict(headers or {})


class BadRequestError(Exception):
    """Same class name as openai.BadRequestError: the harness classifies errors by name (a cell error)."""

    def __init__(self, message):
        super().__init__(message)
        self.response = _HTTPResponse(400)


class RateLimitError(Exception):
    """Same class name as openai.RateLimitError (transient: retried, honouring Retry-After)."""

    def __init__(self, message, retry_after=None):
        super().__init__(message)
        self.response = _HTTPResponse(429, {} if retry_after is None else {"retry-after": retry_after})


class _Completion:
    def __init__(self, payload):
        self._payload = payload

    def model_dump(self):
        return json.loads(json.dumps(self._payload))


class FakeAPI:
    """Deterministic replies for every (model, pair, order) of the sandbox stimuli. round_no 1 is the first
    pass (transient 400s fail); round 2 is the resume (they succeed). One cell's 400 never goes away."""

    P_400, P_429, P_INVALID, P_INVALID_AGAIN, P_THINK_OPEN = 0.0012, 0.0015, 0.02, 0.25, 0.01

    def __init__(self, cases, specs, seed=SEED, round_no=1):
        from r2.analysis.repeated import _snippets
        from r2.analysis.stats import CATEGORY
        from r2.analysis.synthetic import EASE, _behaviour
        from r2.run.harness import ORDERS
        from r2.stimuli.render import render_messages
        self.seed, self.round_no = seed, round_no
        self.lookup = {}
        for cid, case in cases.items():
            snippet = max(_snippets(case) or [""], key=len)
            for order in ORDERS:
                m = render_messages(case, order)
                self.lookup[hashlib.sha256(m["user"].encode("utf-8")).hexdigest()] = (
                    cid, order, m["labels"], EASE[CATEGORY[case["eval"]["sham_trap_type"]]], snippet)
        self.models = {s["model_id"]: (j, s) for j, s in enumerate(specs)}
        self.behaviour = {s["model_id"]: _behaviour(j, len(specs)) for j, s in enumerate(specs)}
        first = sorted(cases)[0]
        self.persistent_cell = (specs[min(1, len(specs) - 1)]["model_id"], first, "ST")
        self.tries = Counter()
        self.events = Counter()

    def factory(self, provider):
        return SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=self.create)),
                               provider=provider)

    def _reply(self, model, cell, retry):
        cid, order, labels, ease, snippet = cell
        spec = self.models[model][1]
        u = lambda *k: _u(self.seed, model, cid, order, retry, *k)  # noqa: E731
        if (not retry and u("invalid") < self.P_INVALID) or (retry and u("invalid2") < self.P_INVALID_AGAIN):
            self.events["invalid_reply"] += 1
            return "I would rely on the first excerpt here." if not retry else "Sorry, I cannot decide."
        if spec.get("reasoning") and not retry and u("think_open") < self.P_THINK_OPEN:
            self.events["unterminated_think"] += 1
            return "<think>\nThe two excerpts look similar; let me compare the dosing line by line"
        p_content, p_first = self.behaviour[model]
        authentic = "A" if labels["A"] == "T" else "B"
        by_content = u("content") < p_content * ease
        sel = authentic if by_content else ("A" if u("first") < p_first else "B")
        other = "B" if sel == "A" else "A"
        v = u("variant")
        raw = f"Excerpt {sel}" if v < 0.05 else (sel.lower() if v < 0.08 else sel)
        conf = 85 if u("conf") < 0.03 else round(0.55 + 0.4 * u("conf2"), 2)
        decision = {"selected_tool": raw, "confidence": conf,
                    "trust_rationale": (f"Excerpt {sel} is consistent with the cited guideline; excerpt {other} "
                                        "was altered." if by_content else
                                        "The two excerpts are essentially identical, so excerpt A was chosen."
                                        if u("rat") < 0.5 else f"Excerpt {sel} gives clearer guidance."),
                    "differences_found": [f"Excerpt {other} contains: {snippet}"] if by_content and snippet else []}
        text = json.dumps(decision)
        if spec.get("reasoning"):
            self.events["think_tags"] += 1
            text = "<think>\nCompare the two excerpts field by field.\n</think>\n" + text
        return text

    def create(self, **kwargs):
        model, messages = kwargs["model"], kwargs["messages"]
        user = next(m["content"] for m in messages if m["role"] == "user")
        cell = self.lookup.get(hashlib.sha256(user.encode("utf-8")).hexdigest())
        retry = len(messages) > 2
        key = (model, cell[0], cell[1]) if cell else (model, "probe", "TS")
        self.events["calls"] += 1
        if key == self.persistent_cell:
            self.events["persistent_400"] += 1
            n = self.events["persistent_400"]
            raise BadRequestError("Error code: 400 - {'error': {'message': 'Invalid prompt: rejected by the provider "
                                  f"safety filter (simulated)', 'request_id': 'req_{n:04d}{'ab' * 8}'}}}}")
        if cell and self.round_no == 1 and _u(self.seed, "400", *key) < self.P_400:
            self.events["transient_400"] += 1
            raise BadRequestError("Error code: 400 - {'error': {'message': 'Upstream model returned an invalid "
                                  "response (simulated)'}}")
        if cell and _u(self.seed, "429", *key) < self.P_429 and self.tries[key] == 0:
            self.tries[key] += 1
            self.events["rate_limited_429"] += 1
            if _u(self.seed, "429kind", *key) < 0.5:
                retry_after = "7"
            else:  # an HTTP-date Retry-After, about 30 s ahead
                retry_after = format_datetime(datetime.now(timezone.utc) + timedelta(seconds=31), usegmt=True)
            raise RateLimitError("Error code: 429 - rate limit exceeded (simulated)", retry_after)
        content = self._reply(model, cell, retry) if cell else json.dumps(
            {"selected_tool": "A", "confidence": 0.7, "trust_rationale": "probe", "differences_found": []})
        return _Completion({"id": f"sim-{self.events['calls']}", "object": "chat.completion", "model": model,
                            "system_fingerprint": None,
                            "choices": [{"index": 0, "finish_reason": "stop",
                                         "message": {"role": "assistant", "content": content}}],
                            "usage": {"prompt_tokens": len(user) // 4, "completion_tokens": len(content) // 4,
                                      "total_tokens": (len(user) + len(content)) // 4}})


class Sleeper:
    """Records the delays the harness asks for (Retry-After, backoff) instead of sleeping."""

    def __init__(self):
        self.delays = []

    def __call__(self, seconds):
        self.delays.append(float(seconds))


# ================================================================ the child: every step in the sandbox
def _child(root, render, seed):
    from r2 import paths, pipeline
    from r2.run.harness import load_specs, load_v2_cases
    from r2.verify import build_packets as BP
    if paths.SANDBOX is None or Path(paths.SANDBOX).resolve() != Path(root).resolve():
        raise SystemExit("the rehearsal child must run with R2_SANDBOX set to the sandbox")
    pipeline.check_write_paths()
    inputs = paths.SANDBOX / "inputs"
    steps, t_all = [], time.time()
    phys = Physicians(seed)
    sleeper = Sleeper()
    api_events = {}

    def step(name, fn, ok_codes=None, expect_refused=False):
        """ok_codes: the result is a command's exit status, and these values count as success. expect_refused:
        the step must be refused (a guard the rehearsal exercises)."""
        t0 = time.time()
        rec = {"step": name}
        try:
            res = fn()
            rec["result"] = res
            rec["status"] = "ok" if ok_codes is None or res in ok_codes else f"exit {res}"
            if expect_refused:
                rec["status"] = "FAILED"
                rec["result"] = f"expected a refusal, got {res!r}"
        except pipeline.NeedsHumans as exc:
            rec.update(status="waiting", result=str(exc))
        except pipeline.Refused as exc:
            rec.update(status="refused (expected)" if expect_refused else "REFUSED", result=str(exc))
        except Exception as exc:  # recorded, then the rehearsal stops
            rec.update(status="FAILED", result=f"{type(exc).__name__}: {exc}")
        rec["seconds"] = round(time.time() - t0, 1)
        steps.append(rec)
        print(f"[{rec['seconds']:7.1f}s] {name}: {rec['status']}", flush=True)
        if rec["status"] in ("REFUSED", "FAILED"):
            raise _Stop(rec)
        return rec

    result = {"steps": steps}
    try:
        step("markers before post-results (pre-results documents)", lambda: pipeline.markers()["counts"])
        before = pipeline.markers()
        queue = json.loads((paths.VERIFY_DIR / "round1" / BP.MANIFEST).read_text())["unresolved"]
        relabelled = relabel_stored_edits(paths.DATA_V2 / "manual_edits.json")
        edits = author_edits_for(queue, paths.DATA_V2 / "manual_edits.json")
        step("simulated round-1 answers and author edits", lambda: {
            "AG": fill_review_workbook(paths.VERIFY_DIR / "round1" / "review_AG.xlsx", inputs / "AG_r1.xlsx", "AG",
                                       phys, author_edits=edits),
            "MO": fill_review_workbook(paths.VERIFY_DIR / "round1" / "review_MO.xlsx", inputs / "MO_r1.xlsx", "MO",
                                       phys),
            "author_edits": len(edits), "stored_edits_relabelled_simulated": relabelled})
        step("pipeline import-edits", lambda: pipeline.import_edits(inputs / "AG_r1.xlsx", SIM_AUTHOR, "2026-10-20"))
        step("pipeline import-edits (rerun: idempotent)",
             lambda: pipeline.import_edits(inputs / "AG_r1.xlsx", SIM_AUTHOR, "2026-10-20"))
        step("simulated round-2 answers", lambda: {
            r: fill_review_workbook(paths.VERIFY_DIR / "round2" / f"review_{r}.xlsx", inputs / f"{r}_r2.xlsx", r, phys)
            for r in ("AG", "MO")})
        books = [inputs / n for n in ("AG_r1.xlsx", "MO_r1.xlsx", "AG_r2.xlsx", "MO_r2.xlsx")]
        step("pipeline consensus (stops for adjudication)", lambda: pipeline.consensus(*books))

        def write_consensus():
            rows = list(csv.DictReader(open(paths.VERIFY_DIR / "adjudication_needed.csv", newline="")))
            with open(inputs / "consensus.csv", "w", newline="") as fh:
                w = csv.writer(fh)
                w.writerow(["case_id", "item", "consensus"])
                for r in rows:
                    w.writerow([r["case_id"], r["item"], phys.truth(r["case_id"], r["item"])])
            return len(rows)
        step("simulated consensus for every disagreement", write_consensus)
        step("pipeline consensus", lambda: pipeline.consensus(*books, consensus_file=inputs / "consensus.csv"))
        step("simulated pilot change: a higher max_tokens for one non-reasoning model", raise_one_max_tokens)
        cases, _ = load_v2_cases()
        specs = load_specs()
        api1 = FakeAPI(cases, specs, seed, round_no=1)
        step("pipeline probe (fake API)", lambda: pipeline.probe(True, api1.factory, sleeper), ok_codes=(0,))
        step("pipeline pilot (fake API, 10 pairs)", lambda: pipeline.run_harness("pilot", True, api1.factory, sleeper),
             ok_codes=(0, 1))
        step("pipeline pilot-summary", lambda: {r["model"]: f"{r['ok']}/{r['cells']} ok"
                                                for r in pipeline.pilot_summary()})
        step("pipeline freeze --dry-run", lambda: pipeline.freeze(dry_run=True))
        step("pipeline freeze --yes (sandbox freeze repository)", lambda: pipeline.freeze(yes=True))
        step("pipeline main before the freeze is public (refused, as designed)",
             lambda: pipeline.run_harness("main", True, api1.factory, sleeper), expect_refused=True)
        step("pipeline record-public (simulated public tag URL, the freeze day)",
             lambda: pipeline.record_public(SIM_TAG_URL, date=_freeze_date(paths.FREEZE_RECORD)))
        step("pipeline main (fake API, first pass)", lambda: pipeline.run_harness("main", True, api1.factory, sleeper),
             ok_codes=(0, 1))
        api2 = FakeAPI(cases, specs, seed, round_no=2)
        step("pipeline main (resume)", lambda: pipeline.run_harness("main", True, api2.factory, sleeper),
             ok_codes=(0,))
        api_events.update({"first_pass": dict(api1.events), "resume": dict(api2.events),
                           "sleeps_requested_s": sleeper.delays, "persistent_cell": list(api1.persistent_cell)})
        step("simulated release tag in the sandbox freeze repository",
             lambda: pipeline._git(paths.FREEZE_REPO, *pipeline.GIT_ID, "tag", SIM_RELEASE, check=True).returncode)
        step("pipeline record --release-tag --zenodo-doi (simulated release)",
             lambda: pipeline.record(release_tag=SIM_RELEASE, zenodo_doi="10.5281/zenodo.0000000"))
        step("pipeline post-results", lambda: _summarize_post(pipeline.post_results(render=render)))
        step("pipeline package-runs (run-log manifest and tarball)", lambda: pipeline.package_runs())
        after = pipeline.markers()
        result["markers_before"] = before
        result["markers_after"] = after
    except _Stop:
        pass
    result["api_events"] = api_events
    result["seconds_total"] = round(time.time() - t_all, 1)
    result["persistent_records"] = _persistent_records(paths.RUNS_V2 / "main")
    (paths.SANDBOX / CHILD_RESULT).write_text(json.dumps(result, indent=1, default=str))
    write_report(paths.SANDBOX, result)
    return 0 if all(s["status"] in OK_STATUS for s in steps) and "markers_after" in result else 1


OK_STATUS = ("ok", "waiting", "refused (expected)")


def raise_one_max_tokens():
    """The pilot change the RUNBOOK allows (a higher max_tokens for a model whose non-responses are truncations),
    simulated in the sandbox model file before the probe: the last non-reasoning model gets twice its limit, so the
    documents report a per-class range and the exception (final review I4). Needs two non-reasoning models."""
    from r2 import paths
    doc = json.loads(paths.MODELS_FILE.read_text())
    nr = [m for m in doc["models"] if not m["reasoning"]]
    if len(nr) < 2:
        return "fewer than two non-reasoning models: no change"
    nr[-1]["max_tokens"] *= 2
    paths.MODELS_FILE.write_text(json.dumps(doc, indent=1) + "\n")
    return {nr[-1]["original_name"]: nr[-1]["max_tokens"]}


class _Stop(Exception):
    pass


def _today():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")


def _freeze_date(record):
    import re
    m = re.search(r"^- Date: (\d{4}-\d{2}-\d{2})$", Path(record).read_text(encoding="utf-8"), re.M)
    return m.group(1) if m else _today()


def _summarize_post(steps):
    out = {}
    for name, rec in steps.items():
        res = rec["result"]
        if isinstance(res, dict) and "tail" in res:
            res = {"rc": res["rc"], "seconds": res["seconds"]}
        elif isinstance(res, list):
            res = [{"rc": r.get("rc"), "command": r.get("command")} if isinstance(r, dict) else r for r in res]
        out[name] = {"seconds": rec["seconds"], "result": res}
    return out


def _persistent_records(main_dir):
    out = []
    for p in sorted(Path(main_dir).glob("*.jsonl")):
        if p.name.startswith("._") or p.name.endswith(".incidents.jsonl"):
            continue
        for line in p.read_bytes().split(b"\n"):
            if b'"persistent_error": true' in line:
                r = json.loads(line)
                out.append({"model": r["original_name"], "case_id": r["case_id"], "order": r["order"],
                            "error_key": r.get("error_key")})
    return out


# ================================================================ report
def _md_cell(x):
    return str(x).replace("|", "\\|").replace("\n", " ")


def write_report(root, result, guard=None):
    root = Path(root)
    L = ["# Rehearsal report (SIMULATED: nothing here is data)", "",
         f"Sandbox: `{root}`. Simulated physicians (seed {SEED}), simulated author edits (author \"{SIM_AUTHOR}\") "
         "and a fake OpenAI-compatible API injected through the harness client factory.", "",
         "## Steps", "", "| # | step | status | seconds | result |", "|---|---|---|---|---|"]
    for i, s in enumerate(result["steps"], 1):
        res = json.dumps(s.get("result"), default=str)
        L.append(f"| {i} | {_md_cell(s['step'])} | {s['status']} | {s['seconds']} | {_md_cell(res[:700])} |")
    L += ["", f"Total: {result.get('seconds_total')} s.", ""]
    ev = result.get("api_events") or {}
    if ev:
        delays = ev.get("sleeps_requested_s", [])
        L += ["## Fake API", "",
              f"- first pass (probe, pilot, main): {json.dumps(ev.get('first_pass'))}",
              f"- resume: {json.dumps(ev.get('resume'))}",
              f"- sleeps requested by the harness (recorded, not slept): {len(delays)}, "
              f"total {sum(delays):.0f} s; distinct values {sorted(set(round(d) for d in delays))}",
              f"- persistent-error cell: {ev.get('persistent_cell')}; persistent_error records in the main run: "
              f"{json.dumps(result.get('persistent_records'))}", ""]
    post = next((s for s in result["steps"] if s["step"] == "pipeline post-results" and s["status"] == "ok"), None)
    if post:
        res = post["result"]
        qa = (res.get("qa") or {}).get("result") or {}
        oc = (res.get("order_check") or {}).get("result") or {}
        L += ["## Post-results checks", ""]
        L += [f"- QA {s}: {v.get('failures')} failure(s), exit status {v.get('rc')}" for s, v in qa.items()]
        L += [f"- order check: freeze {oc.get('freeze_date')}, public since {oc.get('public_since')}, first "
              f"main-run call {oc.get('first_main_call')}: {'PASS' if oc.get('ok') else 'FAIL'}", ""]
    b, a = result.get("markers_before"), result.get("markers_after")
    if a:
        L += ["## Markers", "",
              f"- before post-results (the pre-results documents copied in): {json.dumps(b['counts'])}",
              f"- after post-results: {json.dumps(a['counts'])}",
              f"- NUMERIC markers filled: {b['counts']['NUMERIC'] - a['counts']['NUMERIC']} of "
              f"{b['counts']['NUMERIC']}; NOT filled: {a['counts']['NUMERIC']} (goal 0)", "",
              "### NUMERIC markers not filled", ""]
        L += [f"- {m['document']}: {m['marker']}" for m in a["NUMERIC"]] or ["- none"]
        L += ["", "### UNCLASSIFIED markers", ""]
        L += [f"- {m['document']}: {m['marker']}" for m in a["UNCLASSIFIED"]] or ["- none"]
        L += ["", "### PROSE markers remaining (author-written results text)", ""]
        L += [f"- {m['document']}: {m['marker']}" for m in a["PROSE"]] or ["- none"]
        L.append("")
    if guard is not None:
        L += ["## Real-path check", "",
              f"- unchanged: **{guard['unchanged']}** ({guard['n_files']} real files hashed before and after, "
              f"{guard['n_docx']} of them .docx; git HEAD, tags and status of each checkout compared)",
              f"- roots: {', '.join(guard['roots'])}",
              f"- changed: {guard['changed'] or 'none'}; added: {guard['added'] or 'none'}; removed: "
              f"{guard['removed'] or 'none'}; git changed: {guard['git_changed'] or 'none'}", ""]
    (root / REPORT).write_text("\n".join(L), encoding="utf-8")


# ================================================================ the parent
def child_env(root):
    """The child's environment: R2_SANDBOX=root, R2_REPO=this checkout, and no GIT_* variable
    (r2.sandbox_git.clean_env), so no git command of the child can be redirected into a real repository."""
    from r2 import paths
    from r2.sandbox_git import clean_env
    env = clean_env({**os.environ, "R2_SANDBOX": str(root), "R2_REPO": str(paths.REPO),
                     "PYTHONDONTWRITEBYTECODE": "1"})
    env.pop("R2_RUNS_V2", None)
    return env


def rehearse(root, n_models=None, render=True, seed=SEED, timeout=None):
    """Run the rehearsal in a new sandbox `root`; returns (exit status, guard). The real state must be unchanged."""
    from r2 import paths, pipeline
    from r2.analysis.synthetic import assert_sandbox
    if paths.SANDBOX is not None:
        raise SystemExit("run r2.rehearse without R2_SANDBOX: the rehearsal sets it for its own child process")
    root = assert_sandbox(root)
    before = real_state()
    pipeline.sandbox_init(root, n_models)
    env = child_env(root)
    cmd = [sys.executable, "-m", "r2.rehearse", "--child", str(root), "--seed", str(seed)]
    if not render:
        cmd.append("--no-render")
    with open(root / "rehearsal.log", "w") as log:
        rc = subprocess.run(cmd, cwd=paths.REPO, env=env, stdout=log, stderr=subprocess.STDOUT,
                            timeout=timeout).returncode
    guard = compare_state(before, real_state())
    result = json.loads((root / CHILD_RESULT).read_text()) if (root / CHILD_RESULT).exists() else {"steps": []}
    write_report(root, result, guard)
    if not guard["unchanged"]:
        rc = rc or 4
    return rc, guard


def main(argv=None):
    ap = argparse.ArgumentParser(prog="python -m r2.rehearse", description=__doc__.splitlines()[0])
    ap.add_argument("dir")
    ap.add_argument("--models", type=int, default=None, help="cut the sandbox model list to N (keeps the S5 model)")
    ap.add_argument("--no-render", action="store_true", help="skip the LibreOffice PDF renders")
    ap.add_argument("--seed", type=int, default=SEED)
    ap.add_argument("--child", action="store_true", help=argparse.SUPPRESS)
    a = ap.parse_args(argv)
    if a.child:
        return _child(Path(a.dir), not a.no_render, a.seed)
    rc, guard = rehearse(a.dir, a.models, not a.no_render, a.seed)
    print(f"rehearsal exit {rc}; real state unchanged: {guard['unchanged']}; report: {Path(a.dir) / REPORT}")
    return rc


if __name__ == "__main__":
    sys.exit(main())
